"""AgentOS app construction + export download endpoint (P3b exit criteria)."""

import asyncio
import json
import re

import pytest

from rcm_runbook.config import settings
from rcm_runbook.export.excel import export_xlsx
from tests.unit.test_compliance import full_session


class TestApp:
    def test_agentos_routes_mounted(self, client, con_llave):
        # Con `con_llave` y no con «la llave si la hay»: en un entorno sin .env
        # el test pasaba sin ejercer el cierre de la API, que es justo lo que
        # dice comprobar.
        schema = client.get("/openapi.json", params={"key": con_llave}).json()
        paths = list(schema["paths"])
        assert any("agent" in p or "run" in p or "session" in p for p in paths), paths

    def test_download_serves_golden_export(self, client, con_llave, monkeypatch):
        # El definitivo exige que la sesión exista y esté completa: el archivo en
        # disco es una foto, el estado de la sesión es la verdad.
        from rcm_runbook import app as app_module

        sesion = full_session()
        estado = {"session_state": {"rcm": sesion.model_dump(mode="json")}}
        monkeypatch.setattr(
            app_module.agent.db, "get_session", lambda **kw: {"session_data": estado}
        )
        path = export_xlsx(sesion, settings.exports_dir, session_id="any-session")
        url = f"/exports/any-session/{path.name}"
        assert client.get(url).status_code == 401  # sin llave → rechazado
        resp = client.get(url, params={"key": con_llave})
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument"
        )
        assert len(resp.content) > 5000

    def test_download_rejects_traversal(self, client, con_llave):
        params = {"key": con_llave}
        traversal = client.get("/exports/s/..%2F..%2Fetc%2Fpasswd", params=params)
        assert traversal.status_code in (400, 404)
        assert client.get("/exports/s/no-existe.xlsx", params=params).status_code == 404
        assert client.get("/exports/s/archivo.txt", params=params).status_code == 400


class TestApiIsClosedByDefault:
    """La llave de AgnoAPISettings dejaba abiertas 22 rutas GET, entre ellas
    /sessions y /sessions/{id}/runs — o sea las transcripciones completas de
    cada análisis, legibles por cualquiera con la URL pública."""

    LEAKY = [
        "/sessions",
        "/sessions/cualquiera",
        "/sessions/cualquiera/runs",
        "/memories",
        "/metrics",
        "/traces",
        "/eval-runs",
        "/registry",
        "/info",
        "/user_memory_stats",
        "/trace_session_stats",
    ]

    @pytest.mark.parametrize("path", LEAKY)
    def test_sensitive_routes_require_key(self, client, con_llave, path):
        assert client.get(path).status_code == 401, f"{path} respondió sin llave"

    @pytest.mark.parametrize("path", LEAKY)
    def test_sensitive_routes_open_with_key(self, client, con_llave, path):
        resp = client.get(path, headers={"Authorization": f"Bearer {con_llave}"})
        assert resp.status_code != 401, f"{path} rechazó la llave válida"

    def test_wrong_key_is_rejected(self, client, con_llave):
        assert client.get("/sessions", headers={"Authorization": "Bearer no-es"}).status_code == 401
        assert client.get("/sessions", params={"key": "no-es"}).status_code == 401

    def test_demo_page_and_health_stay_public(self, client, con_llave):
        # El enlace del cliente carga la página sin llave; la página autentica
        # sus propias llamadas. Y el monitoreo necesita /health abierto.
        assert client.get("/demo").status_code == 200
        assert client.get("/health").status_code == 200

    @pytest.mark.parametrize("path", ["/demo/", "/health/", "/favicon.ico"])
    def test_trailing_slash_and_favicon_stay_public(self, client, con_llave, path):
        # Corremos por fuera del TrailingSlashMiddleware de agno: sin normalizar
        # la ruta, el cliente que escribe la URL con `/` final ve un JSON de error
        # en vez del chat, y Safari se lleva un 401 al pedir el favicon.
        assert client.get(path).status_code != 401, f"{path} exigió llave"

    def test_non_ascii_key_is_rejected_not_crashed(self, client, con_llave):
        # `compare_digest` sobre str revienta con no-ASCII: era un 500 con
        # traceback que cualquier anónimo disparaba desde internet.
        assert client.get("/sessions", params={"key": "café"}).status_code == 401
        # En la cabecera llega como bytes y uvicorn la decodifica latin-1
        assert client.get(
            "/sessions", headers={"Authorization": "Bearer ñ".encode("latin-1")}
        ).status_code == 401

    def test_real_preflight_passes_but_bare_options_does_not(self, client, con_llave):
        preflight = client.options(
            "/sessions",
            headers={
                "Origin": "https://os.agno.com",
                "Access-Control-Request-Method": "GET",
            },
        )
        assert preflight.status_code != 401
        assert client.options("/sessions").status_code == 401

    def test_401_carries_cors_headers_for_allowed_origin(self, client, con_llave):
        # Envolvemos al CORSMiddleware: sin esto os.agno.com con la llave mala
        # ve un error opaco de CORS en vez del 401 en español.
        resp = client.get("/sessions", headers={"Origin": "https://os.agno.com"})
        assert resp.status_code == 401
        assert resp.headers.get("access-control-allow-origin") == "https://os.agno.com"

    def test_browser_links_accept_key_query_param(self, client, con_llave):
        # `?key=` debe valer en TODA la API, no solo en /exports: si algún día se
        # exporta OS_SECURITY_KEY al entorno, agno activa su propia dependencia
        # (solo Bearer) y estos enlaces mueren en 401. El contenedor recibe la
        # llave como RCM_OS_SECURITY_KEY justamente para evitarlo.
        assert client.get("/sessions", params={"key": con_llave}).status_code != 401
        assert client.get("/exports/x/y.xlsx", params={"key": con_llave}).status_code != 401

    def test_api_docs_are_not_public(self, client, con_llave):
        # El esquema describe toda la superficie de la API; no se regala.
        assert client.get("/openapi.json").status_code == 401
        assert client.get("/docs").status_code == 401


class TestWebsocketAuth:
    """El middleware `require_key` solo cubre HTTP: el handshake del WebSocket
    queda abierto a propósito, porque os.agno.com se autentica *dentro* del
    protocolo (manda {"action":"authenticate","token":…} ya conectado) y exigir
    la llave en el handshake rompería esa integración.

    Lo que sí debe cumplirse siempre: sin token no se ejecuta ninguna acción.
    Si una versión futura de Agno afloja eso, este test lo caza."""

    @pytest.fixture(autouse=True)
    def _requiere_llave_real(self):
        # agno lee la llave del AgnoAPISettings capturado al importar, así que
        # `con_llave` (que parchea `settings`) no sirve aquí: hace falta la real.
        if not settings.os_security_key:
            pytest.skip("no security key configured in this environment")

    def test_socket_opens_but_demands_auth(self, client):
        with client.websocket_connect("/workflows/ws") as ws:
            saludo = ws.receive_json()
            assert saludo.get("requires_auth") is True

    def test_actions_are_refused_without_token(self, client):
        with client.websocket_connect("/workflows/ws") as ws:
            ws.receive_json()
            ws.send_json({"action": "start_workflow", "workflow_id": "cualquiera"})
            assert ws.receive_json()["event"] == "auth_required"

    def test_wrong_token_is_refused(self, client):
        with client.websocket_connect("/workflows/ws") as ws:
            ws.receive_json()
            ws.send_json({"action": "authenticate", "token": "llave-falsa"})
            assert ws.receive_json()["event"] == "auth_error"

    def test_correct_token_authenticates(self, client):
        with client.websocket_connect("/workflows/ws") as ws:
            ws.receive_json()
            ws.send_json({"action": "authenticate", "token": settings.os_security_key})
            assert ws.receive_json()["event"] == "authenticated"


class TestResumeAfterRestart:
    def test_session_state_survives_rebuild(self):
        """Round-trip through the serialization boundary (what AgentOS persists)."""
        from rcm_runbook.models.session import RCMSession

        session = RCMSession()
        session.scope.tag = "P-201A"
        dumped = session.model_dump(mode="json")
        restored = RCMSession.model_validate(dumped)
        assert restored.scope.tag == "P-201A"
        assert restored.schema_version == 1

    def test_future_schema_version_fails_loud_in_spanish(self):
        from rcm_runbook.agent.tools import _load
        from rcm_runbook.models.session import RCMSession

        class Ctx:
            session_state = {"rcm": {**RCMSession().model_dump(mode="json"),
                                     "schema_version": 99}}

        with pytest.raises(ValueError, match="esquema más nuevo"):
            _load(Ctx())


class TestExportsSurviveEphemeralDisk:
    """En Cloudflare Containers el disco se borra al dormir el contenedor. El
    entregable es función pura del estado de la sesión, así que se regenera
    desde la base en vez de conservarse en disco."""

    def test_regenerates_from_session_state_when_file_is_gone(self, client, tmp_path):
        from rcm_runbook import app as app_module
        from tests.unit.test_compliance import full_session

        sesion = full_session()
        guardado = {"session_data": {"session_state": {"rcm": sesion.model_dump(mode="json")}}}
        original = app_module.agent.db.get_session
        app_module.agent.db.get_session = lambda **kw: guardado
        try:
            tag = sesion.scope.tag
            generado = app_module._regenerar_entregable("s1", f"AMEF_{tag}.xlsx")
            assert generado is not None and generado.is_file()
            assert generado.stat().st_size > 5000
        finally:
            app_module.agent.db.get_session = original

    def test_unknown_session_or_filename_is_not_served(self, client):
        from rcm_runbook import app as app_module

        original = app_module.agent.db.get_session
        app_module.agent.db.get_session = lambda **kw: None
        try:
            assert app_module._regenerar_entregable("no-existe", "AMEF_X.xlsx") is None
        finally:
            app_module.agent.db.get_session = original


class TestDefinitiveExportKeepsItsGate:
    """El nombre del archivo lo elige quien arma la URL. Sin compuerta, quitar el
    prefijo BORRADOR_ entregaba un "definitivo" de un análisis a medias — justo
    lo que la norma prohíbe y lo que la herramienta de exportación sí bloquea."""

    def _con_sesion(self, monkeypatch, sesion):
        from rcm_runbook import app as app_module

        estado = {"session_state": {"rcm": sesion.model_dump(mode="json")}}
        monkeypatch.setattr(
            app_module.agent.db, "get_session", lambda **kw: {"session_data": estado}
        )
        return app_module

    def test_incomplete_session_cannot_yield_a_definitive_file(self, client, monkeypatch):
        from rcm_runbook.engine import compliance
        from rcm_runbook.models.session import RCMSession

        sesion = RCMSession()
        sesion.scope.tag = "P-999"
        assert compliance.export_blockers(sesion), "la sesión debía estar incompleta"
        app_module = self._con_sesion(monkeypatch, sesion)
        assert app_module._regenerar_entregable("s", "AMEF_P-999.xlsx") is None

    def test_incomplete_session_still_yields_a_draft(self, client, monkeypatch):
        from rcm_runbook.models.session import RCMSession

        sesion = RCMSession()
        sesion.scope.tag = "P-999"
        app_module = self._con_sesion(monkeypatch, sesion)
        borrador = app_module._regenerar_entregable("s", "BORRADOR_AMEF_P-999.xlsx")
        assert borrador is not None and borrador.is_file()

    def test_complete_session_yields_the_definitive_file(self, client, monkeypatch):
        from tests.unit.test_compliance import full_session

        sesion = full_session()
        app_module = self._con_sesion(monkeypatch, sesion)
        definitivo = app_module._regenerar_entregable("s", f"AMEF_{sesion.scope.tag}.xlsx")
        assert definitivo is not None and definitivo.is_file()


class TestExportNamesAreSane:
    def test_dot_names_are_rejected(self, client, con_llave):
        # `\w` dejaba pasar «.» y «..» como session_id. El cliente HTTP ya
        # normaliza esas rutas, así que aquí se comprueba el guardia directo.
        from rcm_runbook.app import _SAFE_NAME

        for malo in (".", "..", "..."):
            assert not _SAFE_NAME.match(malo), f"{malo} no fue rechazado"
        for bueno in ("demo-abc", "P-101", "a.b"):
            assert _SAFE_NAME.match(bueno), f"{bueno} fue rechazado por error"

    def test_favicon_is_served_not_404(self, client):
        r = client.get("/favicon.ico")
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("image/svg")


class TestDatabaseBackendIsConfigurable:
    """Sin `DATABASE_URL` se usa SQLite (desarrollo). Con ella, Postgres — que es
    lo que sostiene la continuidad de sesión donde el disco no persiste."""

    def test_sqlite_by_default(self):
        from agno.db.sqlite import SqliteDb

        from rcm_runbook.agent.factory import build_db
        from rcm_runbook.config import Settings

        assert isinstance(build_db(Settings(db_url="")), SqliteDb)

    def test_postgres_when_database_url_is_set(self, monkeypatch):
        from rcm_runbook.agent import factory

        creada = {}

        class FalsaPostgresDb:
            def __init__(self, db_url):
                creada["url"] = db_url

        import agno.db.postgres as pg

        monkeypatch.setattr(pg, "PostgresDb", FalsaPostgresDb)
        from rcm_runbook.config import Settings

        factory.build_db(Settings(db_url="postgresql+psycopg://u:p@h/d"))
        assert creada["url"] == "postgresql+psycopg://u:p@h/d"


class TestCurrentExportRoute:
    """`/exports/{sesion}` sin nombre de archivo: lo que usa el botón de la
    página, porque el navegador no conoce el TAG del equipo."""

    def _con_sesion(self, monkeypatch, sesion):
        from rcm_runbook import app as app_module

        estado = {"session_state": {"rcm": sesion.model_dump(mode="json")}}
        monkeypatch.setattr(
            app_module.agent.db, "get_session", lambda **kw: {"session_data": estado}
        )

    def test_complete_session_gets_the_definitive_file(self, client, con_llave, monkeypatch):
        from tests.unit.test_compliance import full_session

        sesion = full_session()
        self._con_sesion(monkeypatch, sesion)
        r = client.get("/exports/s1", params={"key": con_llave})
        assert r.status_code == 200
        assert f'filename="AMEF_{sesion.scope.tag}.xlsx"' in r.headers["content-disposition"]
        assert len(r.content) > 5000

    def test_incomplete_session_gets_a_clearly_marked_draft(self, client, con_llave, monkeypatch):
        from rcm_runbook.models.session import RCMSession

        sesion = RCMSession()
        sesion.scope.tag = "P-999"
        self._con_sesion(monkeypatch, sesion)
        r = client.get("/exports/s1", params={"key": con_llave})
        assert r.status_code == 200
        assert 'filename="BORRADOR_AMEF_P-999.xlsx"' in r.headers["content-disposition"]

    def test_session_without_tag_explains_instead_of_failing(self, client, con_llave, monkeypatch):
        from rcm_runbook.models.session import RCMSession

        self._con_sesion(monkeypatch, RCMSession())
        r = client.get("/exports/s1", params={"key": con_llave})
        assert r.status_code == 404
        assert "nada que exportar" in r.json()["detail"]

    def test_requires_the_key(self, client, con_llave):
        assert client.get("/exports/s1").status_code == 401


class TestStaleDefinitiveIsNotServedFromDisk:
    """El archivo en disco es una foto del pasado. Si la sesión retrocedió de
    fase, ese `AMEF_` definitivo ya no debe entregarse: manda el estado actual."""

    def test_disk_file_is_ignored_when_session_became_incomplete(
        self, client, con_llave, monkeypatch, tmp_path
    ):
        from rcm_runbook import app as app_module
        from rcm_runbook.export.excel import export_xlsx
        from rcm_runbook.models.session import RCMSession
        from tests.unit.test_compliance import full_session

        completa = full_session()
        app_module.settings.exports_dir = str(tmp_path)
        escrito = export_xlsx(completa, tmp_path, session_id="s-vieja")

        # la sesión retrocede: ahora está incompleta
        incompleta = RCMSession()
        incompleta.scope.tag = completa.scope.tag
        estado = {"session_state": {"rcm": incompleta.model_dump(mode="json")}}
        monkeypatch.setattr(
            app_module.agent.db, "get_session", lambda **kw: {"session_data": estado}
        )
        assert escrito.is_file(), "el definitivo sigue en disco"
        r = client.get(f"/exports/s-vieja/{escrito.name}", params={"key": con_llave})
        assert r.status_code == 404, "sirvió un definitivo obsoleto desde disco"


class TestTempFilesAreCleanedUp:
    """Cada regeneración crea un mkdtemp; sin limpiarlo, cada clic del botón
    dejaba ~11 KB permanentes en el disco del contenedor."""

    def test_regenerated_download_removes_its_temp_dir(self, client, con_llave, monkeypatch):
        import glob

        from rcm_runbook import app as app_module
        from rcm_runbook.models.session import RCMSession

        sesion = RCMSession()
        sesion.scope.tag = "P-777"
        estado = {"session_state": {"rcm": sesion.model_dump(mode="json")}}
        monkeypatch.setattr(
            app_module.agent.db, "get_session", lambda **kw: {"session_data": estado}
        )
        antes = set(glob.glob("/tmp/rcm-export-*"))
        r = client.get("/exports/s-temp", params={"key": con_llave})
        assert r.status_code == 200
        despues = set(glob.glob("/tmp/rcm-export-*"))
        assert despues <= antes, f"quedaron temporales sin borrar: {despues - antes}"


class TestExportsDirEsEscribible:
    """La imagen corría con /app propiedad de root y el proceso como uid 10001,
    así que `export_excel` moría con «Permission denied: 'data'»: la descarga
    por chat estaba rota en producción mientras la del botón funcionaba."""

    def test_el_dockerfile_da_permisos_al_directorio_de_trabajo(self):
        from pathlib import Path

        dockerfile = Path("Dockerfile").read_text(encoding="utf-8")
        assert "chown -R rcm:rcm /app" in dockerfile, (
            "sin esto el proceso no puede crear data/exports y export_excel falla"
        )

    def test_la_herramienta_escribe_donde_apunta_la_config(self, tmp_path, monkeypatch):
        from rcm_runbook.config import settings as cfg
        from rcm_runbook.export.excel import export_xlsx
        from tests.unit.test_compliance import full_session

        monkeypatch.setattr(cfg, "exports_dir", str(tmp_path / "sin-crear"))
        generado = export_xlsx(full_session(), cfg.exports_dir, session_id="s1")
        assert generado.is_file() and generado.stat().st_size > 5000


class TestEnlaceDeExportSinLlave:
    """La herramienta incrustaba la llave en el enlace que ve el cliente, y esa
    transcripción se guarda en Postgres. Además el modelo evitaba repetir un
    enlace con un secreto dentro, así que el cliente se quedaba sin nada que
    pulsar. Ahora el enlace es relativo y la página lo descarga autenticada."""

    def test_el_enlace_no_contiene_la_llave(self, tmp_path, monkeypatch):
        from rcm_runbook.agent import tools as tools_mod
        from rcm_runbook.config import settings as cfg
        from tests.unit.test_compliance import full_session

        monkeypatch.setattr(cfg, "exports_dir", str(tmp_path))
        monkeypatch.setattr(cfg, "os_security_key", "OSK_secreta_de_prueba")

        class Ctx:
            session_id = "s-enlace"
            session_state = {"rcm": full_session().model_dump(mode="json")}

        salida = tools_mod.export_excel.entrypoint(Ctx(), draft=True)
        assert "OSK_secreta_de_prueba" not in salida, "la llave se filtró al chat"
        assert "key=" not in salida
        assert "[Descargar el Excel](/exports/" in salida, salida


class TestCodigoIsoAusenteNoEsFalloTecnico:
    """Un código fuera del catálogo se propagaba como KeyError, y
    `_spanish_errors` lo disfrazaba de «❌ No se pudo completar la operación».
    El agente lo leía como avería del sistema y se inventaba la causa. La
    confusión nacía en la capa de herramientas, no en el modelo."""

    def _ctx(self):
        class Ctx:
            session_id = "s"
            session_state: dict = {}

        return Ctx()

    def test_devuelve_un_mensaje_de_negocio_no_el_banner_de_fallo(self):
        from rcm_runbook.agent import tools as tools_mod

        salida = tools_mod.explain_iso_code.entrypoint(self._ctx(), code="ZZZZ99")
        assert "❌" not in salida, "un código ausente no es un fallo del sistema"
        assert "no está en el catálogo" in salida
        assert "Códigos disponibles" in salida

    def test_un_codigo_valido_sigue_explicandose(self):
        from rcm_runbook.agent import tools as tools_mod
        from rcm_runbook.models.catalogs import fixture

        valido = fixture().menu.iso14224_failure_mode_codes[0].code
        salida = tools_mod.explain_iso_code.entrypoint(self._ctx(), code=valido)
        assert salida.startswith(valido) and "❌" not in salida

    def test_el_enlace_no_expone_la_ruta_del_contenedor(self):
        from rcm_runbook.agent import tools as tools_mod
        from tests.unit.test_compliance import full_session

        class Ctx:
            session_id = "s-ruta"
            session_state = {"rcm": full_session().model_dump(mode="json")}

        salida = tools_mod.export_excel.entrypoint(Ctx(), draft=True)
        assert "data/exports/" not in salida, "expone la ruta interna del contenedor"


class TestGetProgressNoSeUsaComoCompuerta:
    """El validador midió 3 de 3 rechazos inventados del entregable; con la
    instrucción reforzada, 2 de 3. El modelo decide leyendo la salida de
    `get_progress`, no el system prompt, así que el aviso tiene que viajar ahí."""

    def test_la_salida_desmarca_sus_pendientes_como_bloqueadores_del_export(self):
        from rcm_runbook.agent import tools as tools_mod

        class Ctx:
            session_id = "s-prog"
            session_state: dict = {}

        salida = tools_mod.get_progress.entrypoint(Ctx())
        assert "SOLO de la fase actual" in salida
        assert "export_excel" in salida, "no dirige a la herramienta que decide"


class TestUnRechazoDelMetodoNoEsUnaAveria:
    """El agente trata «❌ No se pudo completar la operación» como sistema roto:
    deja de trabajar, avisa de un fallo técnico y se inventa la causa. Varias
    herramientas emitían ese banner —algunas en inglés— ante datos que
    simplemente no valen. Es la diferencia entre «esto se rompió» y «ese dato
    no es el que va»."""

    BANNER = "No se pudo completar la operación"

    def _ctx(self):
        class Ctx:
            session_id = "s-neg"
            session_state: dict = {}

        return Ctx()

    def test_metodo_de_ffi_escrito_en_espanol(self):
        from rcm_runbook.agent import tools as tools_mod

        salida = tools_mod.calculate_ffi.entrypoint(
            self._ctx(), method="disponibilidad", u_fraction=0.02, mtive_hours=8760
        )
        assert self.BANNER not in salida
        assert "availability" in salida, "no dice cuál es el método correcto"
        # Sin el banner pero en inglés seguiría siendo un mensaje del motor
        # colado en una conversación en español (criterio 1 del UAT).
        assert "expected one of" not in salida
        assert "Método de FFI desconocido" in salida

    def test_lista_de_mted_no_numerica(self):
        from rcm_runbook.agent import tools as tools_mod

        salida = tools_mod.calculate_ffi.entrypoint(
            self._ctx(), method="multi_single",
            mted_list_hours="ocho,diez", mtive_hours=8760, mmf_hours=100,
        )
        assert self.BANNER not in salida
        assert "could not convert" not in salida, "fuga del mensaje de Python"

    def test_rango_invalido_se_explica_en_espanol(self):
        from rcm_runbook.agent import tools as tools_mod

        salida = tools_mod.calculate_ffi.entrypoint(
            self._ctx(), method="availability", u_fraction=0.0, mtive_hours=8760
        )
        assert self.BANNER not in salida
        assert "open interval" not in salida, "el mensaje del motor sale en inglés"
        assert "entre 0 y 1" in salida

    def test_tipo_de_funcion_fuera_del_catalogo(self):
        from rcm_runbook.agent import tools as tools_mod

        salida = tools_mod.confirm_no_functions.entrypoint(
            self._ctx(), kind="secundarias_raras"
        )
        assert self.BANNER not in salida
        assert "is not a valid FunctionKind" not in salida
        assert "proteccion" in salida, "no ofrece los tipos válidos"

    def test_un_fallo_tecnico_de_verdad_sigue_llevando_su_banner(self):
        # El contrapeso: si todo dejara de ser técnico, el agente ya no podría
        # distinguir una avería real y la callaría.
        from rcm_runbook.agent import tools as tools_mod

        class Rota:
            session_id = "s"

            @property
            def session_state(self):
                raise RuntimeError("el almacén de sesiones no responde")

        salida = tools_mod.get_progress.entrypoint(Rota())
        assert self.BANNER in salida

    def test_un_dato_invalido_pide_corregir_no_avisar_a_soporte(self):
        from rcm_runbook.agent import tools as tools_mod

        salida = tools_mod.record_task.entrypoint(
            self._ctx(), failure_mode_id="FM-001", description="x",
            frequency="Diaria", duration_hours=1, discipline="mecanica",
        )
        assert self.BANNER not in salida
        assert "corrija y reintente" in salida
        assert "Diario" in salida, "no ofrece la frecuencia válida del catálogo"


class TestElCatalogoIsoLlegaComoLista:
    """La salida de la herramienta se mantiene en lista aunque el renderizador
    ya dibuje tablas: es lo que el modelo copia, y una lista se lee igual de
    bien en el chat y en el Excel exportado."""

    def test_la_herramienta_devuelve_lista_no_tabla(self):
        from rcm_runbook.agent import tools as tools_mod

        class Ctx:
            session_id = "s-iso"
            session_state: dict = {}

        salida = tools_mod.explain_iso_code.entrypoint(Ctx(), code="QQQ1")
        assert "|" not in salida
        assert "- FTS — " in salida


class TestElPlanNoSaleConTareasContradictorias:
    """El primer análisis completo real produjo dos filas para el mismo modo,
    una «1 h, sin paro» y otra «3 h, requiere paro», y la hoja AMEF mostró en
    silencio la equivocada. El guardián prometía «near-duplicate» en su docstring
    y comparaba texto exacto: bastaba con reformular para colarse."""

    def _sesion_con_un_modo(self):
        from tests.unit.test_compliance import full_session

        s = full_session()
        fm = next(iter(s.failure_modes))
        s.tasks.pop(fm, None)
        return s, fm

    def _tarea(self, fm, descripcion, **campos):
        from rcm_runbook.models.domain import MaintenanceTask

        base = {
            "failure_mode_id": fm, "description": descripcion, "frequency": "Mensual",
            "duration_hours": 1.0, "discipline": "Mecánico", "requires_shutdown": False,
        }
        base.update(campos)
        return MaintenanceTask(**base)

    def test_el_reintento_literal_sigue_siendo_idempotente(self):
        s, fm = self._sesion_con_un_modo()
        s.add_task(self._tarea(fm, "Inspeccionar rodamientos"))
        s.add_task(self._tarea(fm, "  inspeccionar  RODAMIENTOS "))
        assert len(s.tasks[fm]) == 1, "el reintento del LLM duplicó la tarea"

    def test_una_reformulacion_no_se_cuela_como_tarea_nueva(self):
        import pytest

        from rcm_runbook.errors import ReglaDeNegocio

        s, fm = self._sesion_con_un_modo()
        s.add_task(self._tarea(fm, "Inspeccionar rodamientos de bomba P-101"))
        with pytest.raises(ReglaDeNegocio) as exc:
            s.add_task(self._tarea(
                fm, "Inspeccionar rodamientos de bomba P-101 (cojinetes motor y bomba)",
                duration_hours=3.0, requires_shutdown=True,
            ))
        assert "casi idéntica" in str(exc.value)
        assert "requiere paro" in str(exc.value), "no muestra en qué se contradicen"
        assert len(s.tasks[fm]) == 1

    def test_misma_descripcion_con_otros_datos_no_se_descarta_en_silencio(self):
        # Es una corrección, no un reintento: quedarse con la primera deja el
        # dato viejo justo cuando alguien intentaba arreglarlo.
        import pytest

        from rcm_runbook.errors import ReglaDeNegocio

        s, fm = self._sesion_con_un_modo()
        s.add_task(self._tarea(fm, "Termografía del tablero"))
        with pytest.raises(ReglaDeNegocio):
            s.add_task(self._tarea(fm, "Termografía del tablero", duration_hours=4.0))

    def test_dos_tareas_de_verdad_distintas_conviven(self):
        # El contrapeso: fusionar de más borraría una tarea legítima del plan.
        s, fm = self._sesion_con_un_modo()
        s.add_task(self._tarea(fm, "Inspeccionar rodamientos"))
        s.add_task(self._tarea(fm, "Lubricar acoples", discipline="Mecánico"))
        s.add_task(self._tarea(fm, "Análisis de vibraciones del motor"))
        assert len(s.tasks[fm]) == 3


class TestLasSiglasVanConSuDefinicion:
    """Tercera y cuarta instancia del mismo defecto: donde la herramienta da una
    sigla pelada, el modelo inventa su significado. Con los códigos ISO se midió
    (5 códigos, 3 corridas, 15/15 mal); con los métodos de FFI, 4 corridas y 4
    significados distintos, ninguno correcto — y ahí lo inventado fija cada
    cuánto se prueba un dispositivo de seguridad."""

    def test_el_rechazo_de_ffi_explica_cada_metodo(self):
        # Las CINCO definiciones, cada una pegada a SU método: la primera
        # versión comprobaba los identificadores y una sola definición suelta,
        # así que intercambiar los significados de multi_single y single_multi
        # pasaba desapercibido.
        from rcm_runbook.agent import tools as tools_mod
        from rcm_runbook.agent.tools import METODOS_FFI

        class Ctx:
            session_id = "s-ffi"
            session_state: dict = {}

        salida = tools_mod.calculate_ffi.entrypoint(Ctx(), method="disponibilidad")
        for metodo, definicion in METODOS_FFI:
            assert f"- {metodo} — {definicion}" in salida, f"{metodo} sin su definición"

    def test_cada_metodo_de_ffi_dice_lo_que_hace_el_motor(self):
        # Contra el motor, no contra sí mismo: intercambiar las definiciones de
        # multi_single y single_multi en la tabla Y en el docstring dejaba la
        # suite en verde, y es justo el error que contamina un cálculo de
        # seguridad.
        from rcm_runbook.agent.tools import METODOS_FFI

        tabla = dict(METODOS_FFI)
        assert "VARIAS funciones protegidas, UN SOLO dispositivo" == tabla["multi_single"]
        assert "UNA función protegida, VARIOS dispositivos redundantes" == (
            tabla["single_multi"]
        )
        # La verdad vive en engine/ffi.py: multi_single recibe una LISTA de
        # Mted (varias funciones protegidas); single_multi recibe n (varios
        # dispositivos).
        import inspect

        from rcm_runbook.engine import ffi

        assert "mted_list" in inspect.signature(ffi.ffi_multi_single).parameters
        assert "n" in inspect.signature(ffi.ffi_single_multi).parameters

    def test_los_parametros_llegan_al_motor_por_el_camino_correcto(self):
        # Intercambiar el enrutado real (mted_list ↔ n) daba otro número con un
        # ✔ delante y ningún test lo miraba.
        from rcm_runbook.agent import tools as tools_mod

        class Ctx:
            session_id = "s-ffi2"
            session_state: dict = {}

        multi = tools_mod.calculate_ffi.entrypoint(
            Ctx(), method="multi_single", mtive_hours=43800,
            mted_list_hours="8760,17520", mmf_hours=100000,
        )
        single = tools_mod.calculate_ffi.entrypoint(
            Ctx(), method="single_multi", mtive_hours=43800, mted_hours=8760,
            mmf_hours=100000, n_devices=2,
        )
        assert "❌" not in multi and "❌" not in single
        # Los dos métodos son fórmulas distintas: si el enrutado se cruzara,
        # darían el mismo número o uno de los dos fallaría.
        assert multi != single

    def test_el_docstring_define_los_parametros_que_se_inventaban(self):
        # Es lo que el modelo lee antes de preguntar. `cff` se leyó como
        # «coeficiente adimensional» y `mtive` como «duración de la prueba».
        from rcm_runbook.agent import tools as tools_mod

        doc = tools_mod.calculate_ffi.entrypoint.__doc__ or ""
        assert "COSTO en dinero de ejecutar una búsqueda" in doc
        assert "TPEF del DISPOSITIVO DE PROTECCIÓN" in doc
        assert "No es la duración de la prueba" in doc

    def test_la_politica_de_mantenimiento_no_llega_como_acronimo_pelado(self):
        # Ejecutando la herramienta, no leyendo su código: la primera versión de
        # este test miraba el fuente y sobrevivía a la mutación que devolvía el
        # acrónimo pelado.
        from rcm_runbook.agent import tools as tools_mod
        from tests.unit.test_compliance import full_session

        sesion = full_session()
        # El modo OCULTO: el FFI solo aplica a fallas ocultas de dispositivos de
        # protección, y la herramienta lo rechaza para los evidentes.
        fm = next(f for f in sesion.failure_modes if sesion.effects[f].is_hidden)

        class Ctx:
            session_id = "s-pol"
            session_state = {"rcm": sesion.model_dump(mode="json")}

        salida = tools_mod.run_decision_logic.entrypoint(
            Ctx(), failure_mode_id=fm, pf_interval_sufficient=True,
            approver="Ana Pérez, Supervisora HSE",
        )
        from rcm_runbook.models.catalogs import POLICY_LABELS_ES

        # Contra el nombre real, no contra «hay un paréntesis»: la primera
        # versión de esta aserción se conformaba con el «(» de «(consecuencia:»
        # y sobrevivía a la mutación que devolvía el acrónimo pelado.
        assert any(n in salida for n in POLICY_LABELS_ES.values()), (
            f"la política llegó como acrónimo pelado: {salida[:100]}"
        )

    def test_sin_politica_determinable_no_es_una_averia(self):
        # El motor está diciendo qué falta preguntarle al equipo. Con el banner
        # técnico delante, el agente lo leía como sistema roto y se paraba.
        from rcm_runbook.agent import tools as tools_mod
        from tests.unit.test_compliance import full_session

        sesion = full_session()
        # El modo EVIDENTE a propósito: en el oculto con consecuencia de
        # seguridad el motor pide confirmación humana antes de decidir, y lo que
        # se comprueba aquí es otra cosa.
        fm = next(f for f in sesion.failure_modes if not sesion.effects[f].is_hidden)

        class Ctx:
            session_id = "s-nopol"
            session_state = {"rcm": sesion.model_dump(mode="json")}

        salida = tools_mod.run_decision_logic.entrypoint(Ctx(), failure_mode_id=fm)
        assert "No se pudo completar la operación" not in salida
        assert "No hay política determinable" in salida


def _registrar_integrante(tools_mod, ctx, nombre: str) -> str:
    """Fuera del bucle a propósito: un lambda que captura la variable de la
    iteración es justo la clase de cierre que ya causó la fuga de la llave."""
    return tools_mod.record_team_member.entrypoint(ctx, name=nombre, role="Operador")


class TestNoSePierdenEscriturasEnParalelo:
    """En producción `score_risk` contestó «✔ Riesgo inicial de FM-002:
    S9-O3-D8 → RPN=216» y seis turnos después el entregable se bloqueó porque
    FM-002 no tenía valoración: el dato de riesgo de un modo con S=9 se evaporó
    con un ✔ delante.

    `_load` lee el estado, la herramienta muta una copia y `_save` reescribe el
    objeto entero. Con dos llamadas del mismo turno en paralelo, la segunda
    parte de la foto vieja y borra a la primera.

    La carrera NO se reproduce sola: en local el load-mutate-save dura
    microsegundos y las hebras casi nunca se solapan — probado, con 24 hebras y
    el candado quitado no se perdía ni una. Por eso se fuerza el solapamiento
    con una barrera: las hebras se esperan unas a otras DENTRO de la lectura,
    que es justo la ventana del defecto. Con el candado puesto no pueden
    coincidir ahí, la barrera vence por tiempo y cada una escribe en su turno.
    """

    HEBRAS = 4

    def _contexto(self):
        import threading

        barrera = threading.Barrier(self.HEBRAS, timeout=0.15)

        class EstadoConBarrera(dict):
            def get(self, clave, defecto=None):
                # La espera va DESPUÉS de leer: así todas se llevan la misma
                # foto vieja y la carrera es segura, no cuestión de suerte.
                # Con la barrera antes de la lectura el defecto solo aparecía
                # 1 de cada 5 veces, porque una hebra terminaba su turno entero
                # antes de que la siguiente llegara a leer.
                valor = super().get(clave, defecto)
                try:
                    barrera.wait()
                except threading.BrokenBarrierError:
                    pass  # el candado hizo su trabajo: no pudieron coincidir
                return valor

        class Ctx:
            session_id = "s-carrera"
            session_state = EstadoConBarrera()

        return Ctx()

    def test_registros_simultaneos_no_se_pisan(self):
        # Cinco rondas: con una sola, quitar el candado fallaba 2 de 3 veces, y
        # un test que deja pasar el defecto un tercio de las veces no sirve de
        # red. Con cinco, que las cinco se salven por casualidad es despreciable.
        import concurrent.futures as cf
        import functools

        from rcm_runbook.agent import tools as tools_mod
        from rcm_runbook.models.session import RCMSession

        for ronda in range(5):
            ctx = self._contexto()
            nombres = [f"Persona {ronda}-{i}" for i in range(self.HEBRAS)]
            with cf.ThreadPoolExecutor(max_workers=self.HEBRAS) as pool:
                list(pool.map(
                    functools.partial(_registrar_integrante, tools_mod, ctx), nombres
                ))

            guardada = RCMSession.model_validate(ctx.session_state["rcm"])
            perdidos = set(nombres) - {m.name for m in guardada.team}
            assert not perdidos, (
                f"ronda {ronda}: se perdieron {len(perdidos)} de {len(nombres)} "
                f"escrituras: {sorted(perdidos)}"
            )

    def test_el_candado_es_por_sesion_no_global(self):
        # Serializar TODAS las sesiones convertiría el candado en un cuello de
        # botella para todos los clientes a la vez.
        from rcm_runbook.agent.tools import _candado

        class A:
            session_id = "sesion-a"

        class B:
            session_id = "sesion-b"

        assert _candado(A()) is not _candado(B())
        assert _candado(A()) is _candado(A())


class TestElNucleoDeLaTareaMandaSobreLosArticulos:
    """Se coló hasta el Excel del cliente quitando dos artículos: «…de los
    rodamientos de la bomba P-101» frente a «…de rodamientos de bomba P-101».
    Y a la vez rechazaba de más: «Análisis de aceite» se tomaba por «Análisis de
    vibración»."""

    BASE = "Análisis de vibración de los rodamientos de la bomba P-101"

    def _casi(self, otro: str) -> bool:
        from rcm_runbook.models.session import RCMSession

        return RCMSession._casi_igual(self.BASE, otro)

    def test_quitar_articulos_no_cuela_una_tarea_nueva(self):
        assert self._casi("Análisis de vibración de rodamientos de bomba P-101")

    def test_el_plural_no_cuela_una_tarea_nueva(self):
        assert self._casi("Análisis de vibraciones de los rodamientos de bomba P-101")

    def test_una_cola_de_relleno_no_desactiva_el_guardian(self):
        # Añadir veinte palabras idénticas diluía el Jaccard bajo el umbral.
        assert self._casi(self.BASE + " durante el turno de mañana con el equipo "
                          "de predictivo presente y su registro correspondiente")

    def test_una_tarea_legitimamente_distinta_no_se_rechaza(self):
        for otra in (
            "Análisis de aceite de los rodamientos de la bomba P-101",
            "Lubricación de los rodamientos de la bomba P-101",
            "Termografía del tablero eléctrico",
            "Alineación láser del acople de la bomba P-101",
        ):
            assert not self._casi(otra), f"rechazó de más: {otra}"


class TestElPlanContradictorioNoSePuedeExportar:
    """La red que sí es determinista. El guardián de texto no puede cazar un
    sinónimo sin rechazar tareas legítimas; esto no adivina si son la misma
    tarea, solo se niega a entregar un plan donde dos filas del mismo modo
    mandan cosas distintas."""

    def _con_dos_tareas(self, **segunda):
        from rcm_runbook.models.domain import MaintenanceTask
        from tests.unit.test_compliance import full_session

        s = full_session()
        fm = next(iter(s.failure_modes))
        s.tasks[fm] = []
        base = {
            "failure_mode_id": fm, "frequency": "Mensual", "duration_hours": 2.0,
            "discipline": "Mecánico", "requires_shutdown": False,
        }
        s.tasks[fm].append(MaintenanceTask(
            description="Análisis de vibración de los rodamientos de la bomba P-101",
            **base,
        ))
        s.tasks[fm].append(MaintenanceTask(**{**base, **segunda}))
        return s, fm

    def test_dos_tareas_que_se_contradicen_bloquean_el_entregable(self):
        from rcm_runbook.engine import compliance

        s, fm = self._con_dos_tareas(
            description="Análisis de vibraciones de rodamientos de bomba P-101",
            requires_shutdown=True,
            duration_hours=3.0, frequency="Semestral",
        )
        bloqueos = compliance.export_blockers(s)
        contradiccion = [b for b in bloqueos if "se contradicen" in b]
        assert contradiccion, f"el plan contradictorio se exporta: {bloqueos}"
        assert fm in contradiccion[0]
        # Los cuatro campos, no solo el paro: comparar uno solo cazaba el
        # incidente original por casualidad.
        for etiqueta in ("frecuencia", "duración", "requiere paro"):
            assert etiqueta in contradiccion[0], f"no señala {etiqueta}"

    def test_solo_la_frecuencia_distinta_ya_contradice_el_plan(self):
        # El caso que se colaba hasta el .xlsx: sinónimo + dos periodicidades
        # para el mismo trabajo, con el mismo requires_shutdown.
        from rcm_runbook.engine import compliance

        s, _ = self._con_dos_tareas(
            description="Análisis de vibraciones de rodamientos de bomba P-101",
            frequency="Semestral",
        )
        bloqueos = [b for b in compliance.export_blockers(s) if "se contradicen" in b]
        assert bloqueos, "dos periodicidades para el mismo modo llegan al CMMS"
        assert "frecuencia" in bloqueos[0]

    def test_solo_el_ejecutor_distinto_ya_contradice_el_plan(self):
        from rcm_runbook.engine import compliance

        s, _ = self._con_dos_tareas(
            description="Análisis de vibraciones de rodamientos de bomba P-101",
            discipline="Instrumentista",
        )
        assert [b for b in compliance.export_blockers(s) if "ejecutor" in b]

    def test_dos_tareas_compatibles_no_bloquean(self):
        # El contrapeso: un modo puede llevar varias tareas legítimas, y
        # bloquear por tenerlas dejaría el producto inservible.
        from rcm_runbook.engine import compliance

        s, _ = self._con_dos_tareas(
            description="Lubricación de los rodamientos de la bomba P-101",
            frequency="Semestral",
            duration_hours=1.0,
        )
        assert not [b for b in compliance.export_blockers(s) if "se contradicen" in b]


class TestLaLetraDeRutaNoViajaSola:
    """Quinta instancia del patrón. Tres corridas en producción dieron tres
    alfabetos contradictorios: una invirtió las dos familias y glosó «CC» como
    «Consecuencias Catastróficas» (real: Control de Calidad). La letra va al
    entregable del cliente y es la traza auditable de por qué se eligió una
    política, así que inventarla contamina la auditoría."""

    def test_la_ruta_llega_con_su_significado(self):
        from rcm_runbook.agent import tools as tools_mod
        from tests.unit.test_compliance import full_session

        sesion = full_session()
        # El modo OCULTO: el FFI solo aplica a fallas ocultas de dispositivos de
        # protección, y la herramienta lo rechaza para los evidentes.
        fm = next(f for f in sesion.failure_modes if sesion.effects[f].is_hidden)

        class Ctx:
            session_id = "s-ruta"
            session_state = {"rcm": sesion.model_dump(mode="json")}

        salida = tools_mod.run_decision_logic.entrypoint(
            Ctx(), failure_mode_id=fm, pf_interval_sufficient=True,
            approver="Ana Pérez, Supervisora HSE",
        )
        from rcm_runbook.models.catalogs import ROUTE_LABELS_ES

        assert any(n in salida for n in ROUTE_LABELS_ES.values()), (
            f"la ruta llegó como letra pelada: {salida[:120]}"
        )

    def test_estan_las_siete_letras_de_las_dos_familias(self):
        from rcm_runbook.models.catalogs import (
            ROUTE_LABELS_ES,
            EvidentRoute,
            HiddenRoute,
        )

        for r in list(EvidentRoute) + list(HiddenRoute):
            assert r.value in ROUTE_LABELS_ES, f"la ruta {r.value} no tiene significado"

    def test_la_A_es_seguridad_en_ambas_familias(self):
        # Una corrida racionalizó la colisión con «Parecen iguales pero no lo
        # son». Sí lo son: A es seguridad/ambiente en las dos.
        from rcm_runbook.models.catalogs import ROUTE_LABELS_ES

        assert "Seguridad" in ROUTE_LABELS_ES["A"]
        assert "oculta" in ROUTE_LABELS_ES["E"] and "Operacional" in ROUTE_LABELS_ES["E"]
        assert "evidente" in ROUTE_LABELS_ES["B"]


class TestLaCausaNoSeCuelaSinTilde:
    """`cause_must_differ_from_mode` bajaba a minúsculas pero no quitaba
    acentos, así que el modo LITERAL escrito sin tildes pasaba como causa."""

    def _crear(self, causa: str):
        from rcm_runbook.models.domain import FailureMode

        return FailureMode(
            id="FM-001", functional_failure_id="FF-001",
            description="Falla de rodamientos con vibración creciente",
            mechanism="Desgaste abrasivo", iso_code="VIB", cause=causa,
            root_cause="Intervalo de relubricación excedido",
            failure_pattern="Fin de Vida Útil",
        )

    def test_el_modo_sin_tildes_no_pasa_como_causa(self):
        import pytest

        with pytest.raises(ValueError, match="reformulación del modo"):
            self._crear("falla de rodamientos con vibracion creciente")

    def test_el_modo_con_tildes_sigue_sin_pasar(self):
        import pytest

        with pytest.raises(ValueError, match="reformulación del modo"):
            self._crear("Falla de rodamientos con vibración creciente")

    def test_una_causa_de_verdad_sigue_aceptandose(self):
        assert self._crear("Lubricante degradado por contaminación con agua")


class TestElUmbralDeCasiIgualEstaFijado:
    """El número que decide si dos tareas son la misma no tenía un solo assert:
    toda la rama difusa era código sin cubrir, y `_sin_acentos` podía ser la
    identidad con las 316 en verde."""

    def _casi(self, a: str, b: str) -> bool:
        from rcm_runbook.models.session import RCMSession

        return RCMSession._casi_igual(a, b)

    def test_los_acentos_no_distinguen_dos_tareas(self):
        assert self._casi("Termografía del tablero", "Termografia del tablero")

    def test_justo_en_el_umbral_es_la_misma_tarea(self):
        # Jaccard = |comunes| / |unión|. Cuatro palabras comunes y cinco en la
        # unión = 0.8 exacto, que es el umbral. Reordenadas para que no sea una
        # subcadena de la otra, o entraría por el camino del containment y el
        # umbral seguiría sin probarse.
        assert self._casi(
            "inspeccionar rodamiento bomba motor acople",
            "bomba inspeccionar rodamiento motor",
        )

    def test_por_debajo_del_umbral_son_tareas_distintas(self):
        # Tres comunes y cinco en la unión = 0.6: no puede fusionarlas.
        assert not self._casi(
            "inspeccionar rodamiento bomba motor acople",
            "bomba inspeccionar rodamiento",
        )

    def test_una_descripcion_vacia_no_se_parece_a_nada(self):
        assert not self._casi("", "Inspeccionar rodamientos")
        assert not self._casi("   ", "   ")

    def test_los_campos_comparados_son_todos_los_que_cambian_el_plan(self):
        # Reducir la lista a duration_hours devolvía el defecto entero: una
        # corrección de frecuencia, disciplina o paro se tragaba en silencio.
        import inspect

        from rcm_runbook.models.session import RCMSession

        fuente = inspect.getsource(RCMSession.add_task)
        for campo in ("frequency", "duration_hours", "discipline", "requires_shutdown"):
            assert campo in fuente, f"{campo} no se compara: una corrección se perdería"


class TestLaHojaAmefNoCallaTareas:
    """La hoja que el cliente abre primero mostraba `tasks[0]` y callaba el
    resto — en el caso real, callaba justo la corregida. Ninguna heurística de
    texto caza un sinónimo sin rechazar tareas legítimas, así que la garantía
    que sí es absoluta es esta: lo registrado no desaparece del entregable."""

    def _con_dos_tareas(self):
        from rcm_runbook.models.domain import MaintenanceTask
        from tests.unit.test_compliance import full_session

        s = full_session()
        fm = next(iter(s.failure_modes))
        base = {
            "failure_mode_id": fm, "duration_hours": 2.0,
            "discipline": "Mecánico", "requires_shutdown": False,
        }
        s.tasks[fm] = [
            MaintenanceTask(description="Análisis de vibración", frequency="Mensual", **base),
            MaintenanceTask(description="Termografía del motor", frequency="Semestral", **base),
        ]
        return s, fm

    def test_las_dos_tareas_aparecen_en_la_hoja_amef(self):
        from rcm_runbook.export.rows import to_amef_rows

        s, fm = self._con_dos_tareas()
        fila = next(f for f in to_amef_rows(s) if True)
        texto = " ".join(str(v) for v in fila.model_dump().values())
        assert "Análisis de vibración" in texto
        assert "Termografía del motor" in texto, "la hoja AMEF calló una tarea"

    def test_manda_al_plan_en_vez_de_inventar_una_frecuencia(self):
        # Con dos frecuencias distintas, elegir una sería mentir; y dejarla en
        # blanco, esconderlo.
        from rcm_runbook.export.rows import to_amef_rows

        s, _ = self._con_dos_tareas()
        texto = " ".join(str(v) for v in to_amef_rows(s)[0].model_dump().values())
        assert "Ver PLAN DE MANTENIMIENTO" in texto

    def test_con_una_sola_tarea_la_hoja_no_cambia(self):
        from rcm_runbook.export.rows import to_amef_rows

        s, fm = self._con_dos_tareas()
        s.tasks[fm] = s.tasks[fm][:1]
        texto = " ".join(str(v) for v in to_amef_rows(s)[0].model_dump().values())
        assert "Mensual" in texto and "Ver PLAN" not in texto


class TestElGlosarioLlegaAlModelo:
    """`ROUTE_LABELS_ES` arreglaba la salida de `run_decision_logic` pero nunca
    entraba en el contexto del modelo: preguntado a pelo —que es lo que hará el
    cliente al mirar la columna «Falla Evidente (ABCD)» de su Excel— cuatro
    corridas dieron tres alfabetos contradictorios, una de ellas re-asignando
    las siete letras a políticas de mantenimiento."""

    def _instrucciones(self) -> str:
        from rcm_runbook.agent.instructions_es import INSTRUCTIONS_ES

        return INSTRUCTIONS_ES

    def test_las_siete_letras_estan_en_las_instrucciones(self):
        from rcm_runbook.models.catalogs import ROUTE_LABELS_ES

        texto = self._instrucciones()
        for letra, significado in ROUTE_LABELS_ES.items():
            assert f"- {letra} — {significado}" in texto, f"la ruta {letra} no llega"

    def test_las_politicas_estan_en_las_instrucciones(self):
        from rcm_runbook.models.catalogs import POLICY_LABELS_ES

        texto = self._instrucciones()
        for politica, nombre in POLICY_LABELS_ES.items():
            assert f"- {politica.value} — {nombre}" in texto, f"{politica.value} no llega"

    def test_avisa_de_la_colision_de_la_A(self):
        # Una corrida racionalizó la colisión con «Parecen iguales pero no lo
        # son». Sí lo son, y hay que decirlo antes de que la explique.
        assert "AMBAS familias" in self._instrucciones()

    def test_explica_como_corregir_un_dato(self):
        texto = self._instrucciones()
        assert "reemplazar=True" in texto
        assert "record_task" in texto


class TestUnaCorreccionNoSeDescartaEnSilencio:
    """La familia entera. Revertir el arreglo de `add_failure_mode` —volver a
    `return existing`— dejaba las 344 en verde: el arreglo se documentó como
    corregido y no tenía ni un test. Estos cubren los tres mutadores."""

    def _sesion(self):
        from rcm_runbook.models.session import RCMSession

        s = RCMSession()
        s.add_function(kind="primaria", verb="bombear", object="crudo",
                       performance_standard="850 GPM a 150 psi")
        s.add_functional_failure("F-001", "No alcanza el caudal requerido")
        return s

    def _modo(self, s, **cambios):
        base = dict(
            description="Cavitación por NPSH por debajo del requerido",
            mechanism="Cavitación", iso_code="LOO",
            cause="Operación fuera de las condiciones de diseño",
            root_cause="Filtro de succión obstruido",
            failure_pattern="Aleatoria",
        )
        base.update(cambios)
        return s.add_failure_mode("FF-001", **base)

    def test_el_modo_con_otra_causa_no_se_traga_en_silencio(self):
        import pytest

        from rcm_runbook.errors import ReglaDeNegocio

        s = self._sesion()
        self._modo(s)
        with pytest.raises(ReglaDeNegocio, match="reemplazar=True"):
            self._modo(s, cause="Lubricante degradado por agua")
        assert s.failure_modes["FM-001"].cause.startswith("Operación fuera")

    def test_con_reemplazar_la_correccion_del_modo_se_aplica(self):
        s = self._sesion()
        self._modo(s)
        fm = self._modo(s, cause="Lubricante degradado por agua", reemplazar=True)
        assert fm.id == "FM-001" and len(s.failure_modes) == 1
        assert s.failure_modes["FM-001"].cause == "Lubricante degradado por agua"

    def test_un_reintento_que_omite_un_opcional_no_se_toma_por_correccion(self):
        # El diff comparaba el volcado entero: un reintento sin el argumento
        # opcional llegaba con el valor por defecto y se reportaba como
        # «cambia: cause, tpef» cuando nadie tocó el tpef.
        s = self._sesion()
        self._modo(s, tpef_hours=None)
        assert self._modo(s).id == "FM-001", "un reintento parcial disparó el rechazo"

    def test_la_tarea_se_puede_reemplazar_de_verdad(self):
        # El mensaje prometía «la reemplazo» y no había forma de hacerlo: una
        # tarea mal registrada era permanente.
        from rcm_runbook.models.domain import MaintenanceTask

        s = self._sesion()
        self._modo(s)
        base = dict(failure_mode_id="FM-001", discipline="Mecánico",
                    duration_hours=2.0, requires_shutdown=False)
        s.add_task(MaintenanceTask(description="Inspección de rodamientos",
                                   frequency="Mensual", **base))
        s.add_task(MaintenanceTask(description="Inspección de rodamientos",
                                   frequency="Semestral", **base), reemplazar=True)
        assert len(s.tasks["FM-001"]) == 1
        assert s.tasks["FM-001"][0].frequency == "Semestral"

    def test_confirmar_que_no_hay_funciones_primarias_no_es_un_no_op(self):
        # Devolvía «✔ Confirmado» sobre una rama que no existía.
        import pytest

        from rcm_runbook.errors import ReglaDeNegocio
        from rcm_runbook.models.domain import FunctionKind

        s = self._sesion()
        with pytest.raises(ReglaDeNegocio, match="siempre tiene función primaria"):
            s.confirm_no_functions_of_kind(FunctionKind.PRIMARIA)


class TestElEntregableLlevaLoQueLasCompuertasExigen:
    """Cinco datos se entrevistaban al cliente, bloqueaban el entregable hasta
    tenerlos, y luego no llegaban a ninguna hoja. Y el FFI —el número que fija
    cada cuánto se prueba un dispositivo de seguridad— se calculaba, se enseñaba
    con un ✔ y moría en el chat: `DecisionResult.ffi_hours` estaba declarado y
    nadie lo escribía nunca."""

    def _libro(self, sesion):
        import io
        import tempfile

        from openpyxl import load_workbook

        from rcm_runbook.export.excel import export_xlsx

        destino = io.BytesIO()
        with tempfile.TemporaryDirectory() as tmp:
            # export_xlsx recibe el DIRECTORIO y decide el nombre.
            destino.write(export_xlsx(sesion, tmp).read_bytes())
        destino.seek(0)
        return load_workbook(destino)

    def _texto_auditoria(self, sesion) -> str:
        hoja = self._libro(sesion)["AUDITORIA RCM"]
        return " | ".join(
            str(c.value) for fila in hoja.iter_rows() for c in fila if c.value is not None
        )

    def test_el_ffi_calculado_aparece_en_el_entregable(self):
        from rcm_runbook.models.session import FFIRegistro
        from tests.unit.test_compliance import full_session

        s = full_session()
        # El modo OCULTO: el entregable filtra el FFI de los modos que ya no lo
        # necesitan, con el mismo criterio que la compuerta.
        fm = next(f for f in s.failure_modes if s.effects[f].is_hidden)
        s.ffi_por_modo[fm] = FFIRegistro(
            horas=4380.0, metodo="single_single", formula="FFI = 2 * Mtive * Mted / Mmf"
        )
        texto = self._texto_auditoria(s)
        assert "4380" in texto, "el FFI no llega a ninguna celda"
        # El método en español: `single_single` crudo era una sigla pelada más,
        # justo debajo de la fila que existe para definirlas.
        assert "una función protegida, un dispositivo" in texto
        assert "single_single" not in texto
        assert fm in texto

    def test_las_acciones_recomendadas_llegan(self):
        from rcm_runbook.models.domain import RecommendedAction
        from tests.unit.test_compliance import full_session

        s = full_session()
        fm = next(iter(s.failure_modes))
        s.actions[fm] = [RecommendedAction(
            failure_mode_id=fm, what="Instalar sensor de vibración en línea",
            who="Ana Pérez", when="2026-09-30",
            verification="Lectura registrada en el CMMS",
        )]
        texto = self._texto_auditoria(s)
        assert "Instalar sensor de vibración" in texto
        assert "Ana Pérez" in texto and "2026-09-30" in texto

    def test_el_equipo_y_la_gobernanza_llegan(self):
        from tests.unit.test_compliance import full_session

        s = full_session()
        s.review_triggers = ["Cambio de contexto operacional"]
        s.validation_signoff = "Validado por Operaciones — Luis Ramos"
        texto = self._texto_auditoria(s)
        assert s.team[0].name in texto, "el equipo multidisciplinario no llega"
        assert "Cambio de contexto operacional" in texto
        assert "Luis Ramos" in texto

    def test_el_ffi_sin_modo_de_falla_avisa_de_que_no_se_guarda(self):
        from rcm_runbook.agent import tools as tools_mod

        class Ctx:
            session_id = "s-ffi3"
            session_state: dict = {}

        salida = tools_mod.calculate_ffi.entrypoint(
            Ctx(), method="single_single", mtive_hours=43800, mted_hours=8760,
            mmf_hours=100000,
        )
        assert "NO queda registrado" in salida, "se pierde en silencio, como antes"
        assert "failure_mode_id" in salida

    def test_el_ffi_con_modo_de_falla_se_guarda(self):
        from rcm_runbook.agent import tools as tools_mod
        from rcm_runbook.models.session import RCMSession
        from tests.unit.test_compliance import full_session

        sesion = full_session()
        # El modo OCULTO: el FFI solo aplica a fallas ocultas de dispositivos de
        # protección, y la herramienta lo rechaza para los evidentes.
        fm = next(f for f in sesion.failure_modes if sesion.effects[f].is_hidden)

        class Ctx:
            session_id = "s-ffi4"
            session_state = {"rcm": sesion.model_dump(mode="json")}

        ctx = Ctx()
        salida = tools_mod.calculate_ffi.entrypoint(
            ctx, method="single_single", mtive_hours=43800, mted_hours=8760,
            mmf_hours=100000, failure_mode_id=fm,
        )
        assert "sale en el entregable" in salida
        guardada = RCMSession.model_validate(ctx.session_state["rcm"])
        assert fm in guardada.ffi_por_modo
        assert guardada.ffi_por_modo[fm].horas > 0


class TestCorregirNoBorraLoQueNoSeMenciona:
    """El `reemplazar=True` de la ronda 7 sustituía el objeto entero: se
    corrigió una causa y se perdió el TPEF con fuente OREDA, y un descarte
    documentado por no credibilidad —el registro que JA1011 exige conservar— se
    resucitó con la justificación en blanco. Con un ✔ delante las dos veces."""

    BASE = dict(
        functional_failure_id="FF-001",
        description="Cavitación por NPSH por debajo del requerido",
        mechanism="Cavitación", iso_code="LOO",
        cause="Operación fuera de las condiciones de diseño",
        root_cause="Filtro de succión obstruido", failure_pattern="Aleatoria",
    )

    def _ctx(self):
        from rcm_runbook.agent import tools as tools_mod

        class Ctx:
            session_id = "s-corr"
            session_state: dict = {}

        c = Ctx()
        tools_mod.record_function.entrypoint(
            c, kind="primaria", verb="bombear", object="crudo",
            performance_standard="850 GPM a 150 psi",
        )
        tools_mod.record_functional_failure.entrypoint(
            c, function_id="F-001", description="No alcanza el caudal requerido"
        )
        return c

    def _guardado(self, ctx, fmid="FM-001"):
        from rcm_runbook.models.session import RCMSession

        return RCMSession.model_validate(ctx.session_state["rcm"]).failure_modes[fmid]

    def test_corregir_la_causa_no_borra_el_tpef(self):
        from rcm_runbook.agent import tools as tools_mod

        c = self._ctx()
        tools_mod.record_failure_mode.entrypoint(
            c, **self.BASE, tpef_hours=26000, tpef_fuente="OREDA"
        )
        tools_mod.record_failure_mode.entrypoint(
            c, **{**self.BASE, "cause": "Lubricante degradado por agua"}, reemplazar=True
        )
        fm = self._guardado(c)
        assert fm.cause == "Lubricante degradado por agua"
        assert fm.tpef is not None, "se perdió el TPEF al corregir la causa"
        assert fm.tpef.value_hours == 26000

    def test_corregir_no_resucita_un_descarte_documentado(self):
        from rcm_runbook.agent import tools as tools_mod

        c = self._ctx()
        tools_mod.record_failure_mode.entrypoint(
            c, **self.BASE, credible=False,
            non_credible_discard="No hay historial en este contexto operacional",
        )
        tools_mod.record_failure_mode.entrypoint(
            c, **{**self.BASE, "cause": "Otra causa registrada"}, reemplazar=True
        )
        fm = self._guardado(c)
        assert fm.credible is False, "el descarte por no credibilidad se resucitó"
        assert fm.non_credible_discard, "la justificación del descarte quedó vacía"

    def test_reemplazar_sin_nada_que_reemplazar_no_crea_una_entidad_nueva(self):
        # El agente creía corregir y duplicaba: dos periodicidades para el mismo
        # trabajo, con el dato viejo intacto en el entregable.
        from rcm_runbook.agent import tools as tools_mod
        from rcm_runbook.models.session import RCMSession

        c = self._ctx()
        tools_mod.record_failure_mode.entrypoint(c, **self.BASE)
        salida = tools_mod.record_failure_mode.entrypoint(
            c, **{**self.BASE, "description": "Erosión del impulsor por sólidos"},
            reemplazar=True,
        )
        assert "no hay ningún modo con esa descripción" in salida
        guardada = RCMSession.model_validate(c.session_state["rcm"])
        assert len(guardada.failure_modes) == 1, "creó una entidad nueva creyendo corregir"

    def test_un_control_repetido_no_invalida_el_riesgo_ni_la_decision(self):
        # add_control era el único mutador sin guarda, y el único cuyo contenido
        # entra en el hash de obsolescencia: un reintento byte a byte idéntico
        # marcaba como desactualizadas la valoración y la decisión, y en un modo
        # de seguridad obligaba a repetir la firma del supervisor.
        from rcm_runbook.engine import compliance
        from tests.unit.test_compliance import full_session

        s = full_session()
        fm = next(iter(s.failure_modes))
        assert not [b for b in compliance.export_blockers(s) if "desactualizada" in b]
        previo = s.controls[fm][0]
        s.add_control(fm, **{k: v for k, v in previo.model_dump().items()
                             if k != "failure_mode_id"})
        assert len(s.controls[fm]) == 1, "duplicó un control idéntico"
        assert not [b for b in compliance.export_blockers(s) if "desactualizada" in b]


class TestElDocstringNoEnsenaValoresInvalidos:
    """El agente le ofreció al interesado, dos veces en producción, «¿Desgaste?
    ¿Fatiga?» — valores que su propio tool doc le enseñó y que la herramienta
    rechaza. Cada invento cuesta un turno."""

    def test_los_patrones_del_docstring_son_los_del_catalogo(self):
        from rcm_runbook.agent import tools as tools_mod
        from rcm_runbook.models.catalogs import fixture

        doc = tools_mod.record_failure_mode.entrypoint.__doc__ or ""
        validos = fixture().menu.failure_patterns
        for patron in validos:
            assert f"'{patron}'" in doc, f"falta el patrón válido {patron}"
        # 'Desgaste' y 'Fatiga' solo pueden aparecer DESPUÉS del aviso de que
        # no existen, nunca en la lista de valores a usar.
        for inventado in ("Desgaste", "Fatiga"):
            assert doc.index(inventado) > doc.index("No existen"), (
                f"{inventado} se sigue ofreciendo como valor válido"
            )

    def test_avisa_de_que_desgaste_y_fatiga_no_existen(self):
        from rcm_runbook.agent import tools as tools_mod

        doc = " ".join((tools_mod.record_failure_mode.entrypoint.__doc__ or "").split())
        assert "No existen 'Desgaste' ni 'Fatiga'" in doc

    def test_las_frecuencias_del_docstring_son_las_del_catalogo(self):
        from rcm_runbook.agent import tools as tools_mod
        from rcm_runbook.models.catalogs import fixture

        doc = tools_mod.record_task.entrypoint.__doc__ or ""
        for frecuencia in fixture().menu.frequencies:
            assert frecuencia in doc, f"falta la frecuencia {frecuencia}"

    def test_avisa_de_que_quinquenal_es_ambiguo(self):
        from rcm_runbook.agent import tools as tools_mod

        doc = tools_mod.record_task.entrypoint.__doc__ or ""
        assert "'Quinquenal' es ambiguo" in doc
        assert "pregunta al interesado" in doc


class TestLaFaseAvanzaSola:
    """`advance_phase` existía y el modelo no la llamaba: 31 turnos en
    producción, con el interesado pidiéndolo cuatro veces, y la sesión terminó
    en fase 1. Las compuertas P2–P5 no llegaban a ejecutarse nunca y una sesión
    reanudada arrancaba leyendo fase 1.

    No había nada que decidir: `check_gate` ya sabe si la fase está completa y
    `advance_phase` habría aceptado exactamente estos avances."""

    def _ctx(self):
        class Ctx:
            session_id = "s-fase"
            session_state: dict = {}

        return Ctx()

    def _completar_fase_1(self, ctx):
        from rcm_runbook.agent import tools as tools_mod

        tools_mod.record_scope.entrypoint(
            ctx, equipment_family="Bomba Centrífuga", equipment_description="Bomba P-101",
            tag="P-101", location="Planta norte", boundaries="Brida a brida",
            interfaces="Succión y descarga", normal_conditions="24/7",
            objective="Disponibilidad 98%", operating_context="Crudo a 60 °C",
        )
        tools_mod.record_team_member.entrypoint(ctx, name="Ana", role="mantenimiento")
        return tools_mod.record_team_member.entrypoint(ctx, name="Luis", role="operaciones")

    def _fase(self, ctx):
        from rcm_runbook.models.session import RCMSession

        return RCMSession.model_validate(ctx.session_state["rcm"]).phase

    def test_al_completarse_la_compuerta_la_fase_avanza(self):
        from rcm_runbook.models.session import Phase

        ctx = self._ctx()
        self._completar_fase_1(ctx)
        assert self._fase(ctx) == Phase.P2_FUNCIONES, "la fase se quedó congelada"

    def test_la_herramienta_avisa_del_avance(self):
        # Sin el aviso el agente sigue hablando de la fase vieja y repregunta lo
        # que ya está cerrado.
        salida = self._completar_fase_1(self._ctx())
        assert "avanzó a la fase 2" in salida
        assert "Funciones y fallas funcionales" in salida

    def test_no_avanza_con_la_compuerta_en_rojo(self):
        from rcm_runbook.agent import tools as tools_mod
        from rcm_runbook.models.session import Phase

        ctx = self._ctx()
        # Solo un integrante: la compuerta P1 exige mantenimiento + operaciones.
        tools_mod.record_scope.entrypoint(
            ctx, equipment_family="Bomba", equipment_description="P-101", tag="P-101",
            location="Norte", boundaries="Brida a brida", interfaces="Succión",
            normal_conditions="24/7", objective="98%", operating_context="Crudo",
        )
        salida = tools_mod.record_team_member.entrypoint(ctx, name="Ana", role="mantenimiento")
        assert self._fase(ctx) == Phase.P1_ALCANCE
        assert "avanzó a la fase" not in salida

    def test_el_aviso_no_se_repite_en_la_llamada_siguiente(self):
        from rcm_runbook.agent import tools as tools_mod

        ctx = self._ctx()
        self._completar_fase_1(ctx)
        siguiente = tools_mod.get_progress.entrypoint(ctx)
        assert "avanzó a la fase" not in siguiente


class TestElAgenteVeElEstadoSinPedirlo:
    """`add_session_state_to_context=False` dejaba al agente ciego salvo que
    llamara a `get_progress`, y con una ventana de 10 turnos el dato que el
    interesado dio y él no registró salía de la historia y desaparecía. Medido:
    31 turnos, 0 modos de falla, repreguntando cosas que él mismo marcaba con
    «✔ ya lo dijiste»."""

    def test_el_digest_viaja_como_dependencia_del_agente(self):
        import inspect

        from rcm_runbook.agent import factory

        fuente = inspect.getsource(factory.build_agent)
        assert "add_dependencies_to_context=True" in fuente
        assert "_digest_de_la_sesion" in fuente

    def test_el_digest_resume_el_estado_real(self):
        from rcm_runbook.agent.factory import _digest_de_la_sesion
        from tests.unit.test_compliance import full_session

        sesion = full_session()
        texto = _digest_de_la_sesion({"rcm": sesion.model_dump(mode="json")})
        assert "FASE ACTUAL" in texto
        assert sesion.scope.tag in texto
        assert next(iter(sesion.failure_modes)) in texto

    def test_una_sesion_nueva_lo_dice_sin_reventar(self):
        from rcm_runbook.agent.factory import _digest_de_la_sesion

        assert "sin datos" in _digest_de_la_sesion({})
        assert "sin datos" in _digest_de_la_sesion(None)

    def test_un_estado_corrupto_degrada_el_resumen_no_el_turno(self):
        # Perder el resumen empeora la conversación; perder el turno la corta.
        from rcm_runbook.agent.factory import _digest_de_la_sesion

        # Un payload que pydantic sí rechaza: con claves sueltas rellena por
        # defecto y no falla, así que ese caso no probaba nada.
        texto = _digest_de_la_sesion({"rcm": {"phase": "no-es-una-fase"}})
        assert "get_progress" in texto


class TestLaSondaDistingueVivoDeFuncionando:
    """La cuenta del proveedor se quedó sin saldo, cada turno del chat devolvía
    un error, y `/health` respondió 200 durante toda la caída. Un servicio cuya
    única sonda no puede distinguir «funciona» de «no funciona» no es
    observable, por muchos 200 que devuelva."""

    def test_la_sonda_profunda_exige_la_llave(self, client, con_llave):
        # Gasta un token del operador: pública sería una forma cómoda de
        # vaciarle la cuenta a alguien recargando una URL.
        assert client.get("/health/modelo").status_code == 401

    def test_health_sigue_siendo_publico_y_barato(self, client, con_llave):
        assert client.get("/health").status_code == 200

    def test_sin_saldo_responde_degradado_y_lo_explica(self, client, con_llave, monkeypatch):
        from rcm_runbook import app as app_mod

        monkeypatch.setattr(
            app_mod, "sondear_modelo",
            lambda _cfg: (False, "La cuenta del proveedor del modelo no tiene saldo."),
        )
        r = client.get("/health/modelo", headers={"Authorization": f"Bearer {con_llave}"})
        assert r.status_code == 503, "una caída total respondía 200"
        assert r.json()["estado"] == "degradado"
        assert "saldo" in r.json()["detalle"]

    def test_cuando_responde_dice_ok(self, client, con_llave, monkeypatch):
        from rcm_runbook import app as app_mod

        monkeypatch.setattr(
            app_mod, "sondear_modelo", lambda _cfg: (True, "El proveedor responde.")
        )
        r = client.get("/health/modelo", headers={"Authorization": f"Bearer {con_llave}"})
        assert r.status_code == 200 and r.json()["estado"] == "ok"

    def test_el_sondeo_traduce_los_fallos_conocidos(self, monkeypatch):
        from rcm_runbook.agent.factory import sondear_modelo
        from rcm_runbook.config import settings

        class ClienteRoto:
            def __init__(self, *a, **k):
                pass

            class messages:  # noqa: N801
                @staticmethod
                def create(**_kw):
                    raise RuntimeError("Error code: 400 - your credit balance is too low")

        import anthropic

        monkeypatch.setattr(anthropic, "Anthropic", ClienteRoto)
        ok, detalle = sondear_modelo(settings)
        assert ok is False
        assert "saldo" in detalle
        assert "credit balance" not in detalle, "le enseña la facturación a quien sondee"

    def test_un_fallo_que_no_reconoce_tambien_es_degradado(self, monkeypatch):
        # Lo peligroso no es no saber traducirlo: es decir «ok» cuando el
        # proveedor no respondió. Esa es exactamente la forma de la caída que
        # tuvo el producto muerto con la vigilancia en verde.
        from rcm_runbook.agent.factory import sondear_modelo
        from rcm_runbook.config import settings

        class ClienteRoto:
            def __init__(self, *a, **k):
                pass

            class messages:  # noqa: N801
                @staticmethod
                def create(**_kw):
                    raise RuntimeError("algo raro que nadie previó")

        import anthropic

        monkeypatch.setattr(anthropic, "Anthropic", ClienteRoto)
        ok, detalle = sondear_modelo(settings)
        assert ok is False, "un fallo desconocido se reportó como servicio sano"
        assert "no respondió" in detalle


class TestElFfiGobiernaLaFrecuenciaDelCmms:
    """Tres intentos. Comparar TODAS las tareas rechazaba planes correctos;
    comparar «que cumpla alguna» dejaba pasar la prueba atrasada en cuanto
    hubiera una limpieza mensual al lado — el defecto original, reabierto por su
    propio arreglo. El error estaba en deducirlo: el dato no existía en el
    modelo, y ahora existe."""

    def _sesion(self, frecuencia: str, horas_ffi: float = 1752.0, marcada: bool = True):
        from rcm_runbook.models.domain import MaintenanceTask
        from rcm_runbook.models.session import FFIRegistro
        from tests.unit.test_compliance import full_session

        s = full_session()
        # El modo OCULTO: el FFI solo aplica a fallas ocultas de dispositivos de
        # protección, y la compuerta salta los demás a propósito.
        fm = next(
            f for f in s.failure_modes
            if s.effects[f].is_hidden and s.failure_modes[f].credible
        )
        s.ffi_por_modo[fm] = FFIRegistro(
            horas=horas_ffi, metodo="single_single", formula="FFI = 2 * Mtive * Mted / Mmf"
        )
        s.tasks[fm] = [MaintenanceTask(
            failure_mode_id=fm, description="Prueba funcional del disparo",
            frequency=frecuencia, duration_hours=4.0, discipline="Instrumentista",
            requires_shutdown=True, es_busqueda_de_fallas=marcada,
        )]
        return s, fm

    def _otra(self, s, fm, descripcion: str, frecuencia: str):
        from rcm_runbook.models.domain import MaintenanceTask

        s.tasks[fm].append(MaintenanceTask(
            failure_mode_id=fm, description=descripcion, frequency=frecuencia,
            duration_hours=2.0, discipline="Instrumentista",
        ))
        return s

    def _bloqueos(self, s):
        from rcm_runbook.engine import compliance

        return [b for b in compliance.export_blockers(s) if "búsqueda de fallas" in b]

    def test_la_prueba_mas_espaciada_que_el_ffi_bloquea(self):
        s, fm = self._sesion("Semestral")  # 4380 h contra 1752 calculadas
        bloqueos = self._bloqueos(s)
        assert bloqueos and fm in bloqueos[0]
        assert "1752" in bloqueos[0] and "Semestral" in bloqueos[0]

    def test_una_tarea_rapida_al_lado_no_tapa_la_prueba_atrasada(self):
        # El caso que reabrió el defecto: `any` no distingue «existe la prueba a
        # tiempo» de «existe cualquier otra cosa a tiempo».
        s, fm = self._sesion("Semestral")
        self._otra(s, fm, "Inspección visual del tablero", "Diario")
        assert self._bloqueos(s), "una inspección diaria tapó la prueba semestral"

    def test_parada_de_planta_no_se_cuela_ni_con_una_limpieza_mensual(self):
        s, fm = self._sesion("Parada de Planta")
        self._otra(s, fm, "Limpieza del entorno del equipo", "Mensual")
        bloqueos = self._bloqueos(s)
        assert bloqueos, "la prueba a Parada de Planta cruzó la compuerta"
        assert "no tiene equivalencia en horas" in bloqueos[0]

    def test_una_calibracion_anual_no_estorba_si_la_prueba_cumple(self):
        s, fm = self._sesion("Bimestral")  # 1460 h ≤ 1752
        self._otra(s, fm, "Calibración del transmisor", "Anual")
        assert not self._bloqueos(s), "rechazó un plan correcto"

    def test_un_modo_con_ffi_y_sin_tarea_marcada_bloquea(self):
        # `_gate_p6` salta los modos OHF y los no creíbles, así que dar por
        # hecho que ya exigía tareas era falso: exportaba limpio con un
        # intervalo calculado que no ejecutaba nadie.
        s, fm = self._sesion("Bimestral", marcada=False)
        bloqueos = self._bloqueos(s)
        assert bloqueos, "un FFI sin tarea que lo ejecute salió en el entregable"
        assert "ninguna tarea marcada" in bloqueos[0]
        assert "es_busqueda_de_fallas=True" in bloqueos[0], "no dice cómo arreglarlo"

    def test_un_modo_con_ffi_y_cero_tareas_bloquea(self):
        s, fm = self._sesion("Bimestral")
        s.tasks[fm] = []
        assert self._bloqueos(s)

    def test_el_umbral_es_el_intervalo_calculado_no_un_multiplo(self):
        s, _ = self._sesion("Trimestral", horas_ffi=2000.0)  # 2190 > 2000
        assert self._bloqueos(s), "una tarea apenas más lenta que el FFI se coló"
        s, _ = self._sesion("Trimestral", horas_ffi=2190.0)  # exacto
        assert not self._bloqueos(s), "el intervalo exacto debe aceptarse"

    def test_los_valores_del_catalogo_estan_fijados(self):
        # De los 13 valores solo dos estaban sujetos por una prueba: los otros
        # once podían valer cualquier cosa. 'Bi-Anual' son DOS AÑOS —lo dice la
        # ordenación creciente del catálogo, entre Anual y Tri-Anual—, no dos
        # veces al año.
        from rcm_runbook.models.catalogs import FRECUENCIA_EN_HORAS

        esperado = {
            "Diario": 24, "Semanal": 168, "Catorcenal": 336, "Mensual": 730,
            "Bimestral": 1460, "Trimestral": 2190, "Tetramestral": 2920,
            "Semestral": 4380, "Anual": 8760, "Bi-Anual": 17520,
            "Tri-Anual": 26280, "Tetra-Anual": 35040, "Quinque-Annual": 43800,
        }
        assert FRECUENCIA_EN_HORAS == esperado

    def test_el_catalogo_de_horas_esta_ordenado_como_el_del_cliente(self):
        from rcm_runbook.models.catalogs import FRECUENCIA_EN_HORAS, fixture

        del_catalogo = [f for f in fixture().menu.frequencies if f in FRECUENCIA_EN_HORAS]
        horas = [FRECUENCIA_EN_HORAS[f] for f in del_catalogo]
        pares = list(zip(del_catalogo, horas, strict=True))
        assert horas == sorted(horas), f"el orden no crece: {pares}"


class TestElValorDelFfiEstaFijado:
    """Ninguna aserción miraba el VALOR: duplicar el FFI persistido pasaba con
    las 391 en verde, y es el número que fija cada cuánto se prueba un
    dispositivo de seguridad."""

    def test_el_numero_guardado_es_el_que_calcula_el_motor(self):
        from rcm_runbook.agent import tools as tools_mod
        from rcm_runbook.engine.ffi import ffi_single_single
        from rcm_runbook.models.session import RCMSession
        from tests.unit.test_compliance import full_session

        sesion = full_session()
        # El modo OCULTO: el FFI solo aplica a fallas ocultas de dispositivos de
        # protección, y la herramienta lo rechaza para los evidentes.
        fm = next(f for f in sesion.failure_modes if sesion.effects[f].is_hidden)

        class Ctx:
            session_id = "s-valor"
            session_state = {"rcm": sesion.model_dump(mode="json")}

        ctx = Ctx()
        tools_mod.calculate_ffi.entrypoint(
            ctx, method="single_single", mtive_hours=43800, mted_hours=17520,
            mmf_hours=876000, failure_mode_id=fm,
        )
        esperado = ffi_single_single(mtive=43800, mted=17520, mmf=876000).ffi_hours
        guardado = RCMSession.model_validate(ctx.session_state["rcm"]).ffi_por_modo[fm]
        assert guardado.horas == esperado, f"guardó {guardado.horas}, el motor dice {esperado}"
        assert guardado.metodo == "single_single"

    def test_el_numero_del_entregable_es_el_guardado(self):
        import tempfile

        from openpyxl import load_workbook

        from rcm_runbook.export.excel import export_xlsx
        from rcm_runbook.models.session import FFIRegistro
        from tests.unit.test_compliance import full_session

        s = full_session()
        # El modo OCULTO: el entregable filtra el FFI de los modos que ya no lo
        # necesitan, con el mismo criterio que la compuerta.
        fm = next(f for f in s.failure_modes if s.effects[f].is_hidden)
        s.ffi_por_modo[fm] = FFIRegistro(
            horas=1752.0, metodo="single_single", formula="FFI = 2 * Mtive * Mted / Mmf"
        )
        with tempfile.TemporaryDirectory() as tmp:
            aud = load_workbook(export_xlsx(s, tmp))["AUDITORIA RCM"]
        numeros = [c.value for f in aud.iter_rows() for c in f if isinstance(c.value, int | float)]
        assert 1752.0 in numeros, f"el FFI del entregable no es el calculado: {numeros}"


class TestLosHuecosQueDestapoLaMutacion:
    """Cinco mutaciones sobrevivían a las 391 pruebas. No son casos raros: son
    ramas de rechazo sin cobertura y un contrapeso que faltaba."""

    def _sesion(self):
        from rcm_runbook.models.session import RCMSession

        s = RCMSession()
        s.add_function(kind="primaria", verb="bombear", object="crudo",
                       performance_standard="850 GPM a 150 psi")
        s.add_functional_failure("F-001", "No alcanza el caudal requerido")
        s.add_failure_mode(
            "FF-001", description="Cavitación por NPSH por debajo del requerido",
            mechanism="Cavitación", iso_code="LOO",
            cause="Operación fuera de las condiciones de diseño",
            root_cause="Filtro de succión obstruido", failure_pattern="Aleatoria",
        )
        return s

    def test_reemplazar_una_funcion_inexistente_no_la_crea(self):
        import pytest

        from rcm_runbook.errors import ReglaDeNegocio

        s = self._sesion()
        with pytest.raises(ReglaDeNegocio, match="no hay ninguna función"):
            s.add_function(kind="primaria", verb="comprimir", object="gas",
                           performance_standard="500 m³/h", reemplazar=True)
        assert len(s.functions) == 1, "creó una función creyendo corregir"

    def test_reemplazar_una_tarea_inexistente_no_la_crea(self):
        import pytest

        from rcm_runbook.errors import ReglaDeNegocio
        from rcm_runbook.models.domain import MaintenanceTask

        s = self._sesion()
        tarea = MaintenanceTask(
            failure_mode_id="FM-001", description="Termografía del tablero",
            frequency="Mensual", duration_hours=1.0, discipline="Predictivo",
        )
        with pytest.raises(ReglaDeNegocio, match="no hay ninguna tarea parecida"):
            s.add_task(tarea, reemplazar=True)
        assert not s.tasks.get("FM-001"), "añadió una tarea creyendo corregir"

    def test_dos_controles_distintos_conviven(self):
        # El contrapeso de la guarda de idempotencia: si fusionara de más, el
        # análisis perdería controles reales del activo.
        s = self._sesion()
        s.add_control("FM-001", description="Inspección visual mensual", kind="detectivo")
        s.add_control("FM-001", description="Alarma de vibración en DCS", kind="detectivo")
        assert len(s.controls["FM-001"]) == 2

    def test_los_kpis_no_salen_como_repr_de_python(self):
        import tempfile

        from openpyxl import load_workbook

        from rcm_runbook.export.excel import export_xlsx
        from tests.unit.test_compliance import full_session

        s = full_session()
        with tempfile.TemporaryDirectory() as tmp:
            aud = load_workbook(export_xlsx(s, tmp))["AUDITORIA RCM"]
        texto = " | ".join(str(c.value) for f in aud.iter_rows() for c in f if c.value)
        assert s.kpis, "el escenario de prueba no tiene KPIs que comprobar"
        assert "name=" not in texto, "el KPI salió como repr de Python"
        assert s.kpis[0].name in texto, "el KPI no llegó al entregable"

    def test_los_temporales_se_limpian_donde_de_verdad_se_crean(self):
        # El test anterior miraba /tmp/rcm-export-*, y en macOS mkdtemp cae en
        # $TMPDIR: los dos conjuntos eran siempre vacíos y `set() <= set()`
        # siempre cierto.
        import tempfile
        from pathlib import Path

        raiz = Path(tempfile.gettempdir())
        antes = set(raiz.glob("rcm-export-*"))
        from rcm_runbook.export.excel import export_xlsx
        from tests.unit.test_compliance import full_session

        with tempfile.TemporaryDirectory() as tmp:
            export_xlsx(full_session(), tmp)
        assert set(raiz.glob("rcm-export-*")) == antes


class TestNadaEnInglesLlegaAlClienteNiAlModelo:
    """Criterio 1: la conversación es 100% en español. Los avisos del motor de
    FFI salían crudos al chat Y a la columna F del entregable —`_en_espanol`
    envolvía excepciones, no resultados— y el digest que el modelo lee en cada
    turno repetía las siglas peladas mientras el criterio 40 medía si las
    acertaba."""

    def test_el_aviso_del_motor_de_ffi_llega_en_espanol(self):
        from rcm_runbook.agent import tools as tools_mod

        class Ctx:
            session_id = "s-aviso"
            session_state: dict = {}

        salida = tools_mod.calculate_ffi.entrypoint(
            Ctx(), method="single_single", mtive_hours=1, mted_hours=100000, mmf_hours=1
        )
        assert "exceeds the protective device" not in salida
        assert "supera el TPEF del propio dispositivo" in salida

    def test_el_aviso_guardado_en_el_entregable_tambien(self):
        from rcm_runbook.agent import tools as tools_mod
        from rcm_runbook.models.session import RCMSession
        from tests.unit.test_compliance import full_session

        sesion = full_session()
        # El modo OCULTO: el FFI solo aplica a fallas ocultas de dispositivos de
        # protección, y la herramienta lo rechaza para los evidentes.
        fm = next(f for f in sesion.failure_modes if sesion.effects[f].is_hidden)

        class Ctx:
            session_id = "s-aviso2"
            session_state = {"rcm": sesion.model_dump(mode="json")}

        ctx = Ctx()
        tools_mod.calculate_ffi.entrypoint(
            ctx, method="single_single", mtive_hours=1, mted_hours=100000,
            mmf_hours=1, failure_mode_id=fm,
        )
        guardado = RCMSession.model_validate(ctx.session_state["rcm"]).ffi_por_modo[fm]
        assert guardado.avisos, "se perdieron los avisos"
        assert not any("exceeds" in a for a in guardado.avisos), (
            "el entregable lleva el aviso del motor en inglés"
        )

    def test_el_digest_glosa_la_politica(self):
        from rcm_runbook.models.catalogs import POLICY_LABELS_ES
        from tests.unit.test_compliance import full_session

        texto = full_session().digest_es()
        assert any(n in texto for n in POLICY_LABELS_ES.values()), (
            "el digest repite las siglas peladas en cada turno"
        )

    def test_el_digest_no_abrevia_el_tipo_de_funcion(self):
        from tests.unit.test_compliance import full_session

        texto = full_session().digest_es()
        assert "[prot]" not in texto and "[prim]" not in texto
        # «de protección», no «(proteccion)»: al quitar la abreviatura quedó el
        # identificador del enum, sin tilde — una abreviatura cambiada por un
        # identificador mal escrito.
        assert "(de protección)" in texto
        assert "(proteccion)" not in texto

    def test_un_analisis_completado_no_dice_fase_7_de_6(self):
        from rcm_runbook.models.session import Phase
        from tests.unit.test_compliance import full_session

        s = full_session()
        s.phase = Phase.COMPLETADO
        texto = s.digest_es()
        assert "7/6" not in texto, "el modelo lee «fase 7 de 6» al cerrar el análisis"
        assert texto.splitlines()[0] == "FASE ACTUAL: Análisis completado"


class TestLasPruebasCorrenEnUnEntornoLimpio:
    """61 pruebas no llegaban a ejecutarse sin credencial del proveedor —el
    cierre de la API, las rutas de export, la página de demo y la sonda solo
    existían en la máquina del desarrollador— y otras 4 se saltaban en verde sin
    llave. Una suite que desaparece justo donde se comprueba la seguridad no es
    una red.

    Se comprueba el EFECTO, no el texto del conftest: la primera versión de esto
    hacía grep sobre el fuente y fallaba por una palabra dentro de un comentario.
    """

    def test_hay_credencial_del_proveedor_para_poder_importar_la_app(self):
        # Contra `settings`, no contra el entorno: en modo suscripción,
        # `build_model` BORRA ANTHROPIC_API_KEY del proceso a propósito (la API
        # rechaza las peticiones que llevan las dos credenciales), así que
        # mirarlo en os.environ da un falso negativo.
        import os

        from rcm_runbook.config import settings

        assert settings.claude_code_oauth_token or os.environ.get("ANTHROPIC_API_KEY"), (
            "sin credencial la app hace SystemExit y 61 pruebas no llegan a correr"
        )

    def test_hay_llave_configurada_para_que_nada_se_salte_en_verde(self):
        from rcm_runbook.config import settings

        assert settings.os_security_key, (
            "sin llave, las cuatro pruebas del WebSocket se saltan justo donde se "
            "comprueba que el socket exige autenticación"
        )

    def test_la_llave_de_prueba_no_activa_la_auth_propia_de_agno(self, client, con_llave):
        # OS_SECURITY_KEY en el entorno activa la dependencia de agno (solo
        # Bearer) y mata los enlaces `?key=`. Es un defecto que ya llegó a
        # producción, y al escribir la fixture lo reintroduje sin darme cuenta.
        assert client.get("/sessions", params={"key": con_llave}).status_code != 401


class TestLoQueLaRonda10DestapoSinCobertura:
    """Seis mutaciones sobrevivían a las 416: el KPI sin meta, la fila de
    nomenclatura entera, la glosa del digest aceptando cualquier etiqueta, el
    aviso traducido sin su número, el segundo aviso del motor sin traducir, y el
    anuncio de «fase 7 de 6»."""

    def _auditoria(self, sesion) -> str:
        import tempfile

        from openpyxl import load_workbook

        from rcm_runbook.export.excel import export_xlsx

        with tempfile.TemporaryDirectory() as tmp:
            hoja = load_workbook(export_xlsx(sesion, tmp))["AUDITORIA RCM"]
        return " | ".join(
            str(c.value) for f in hoja.iter_rows() for c in f if c.value is not None
        )

    def test_el_kpi_llega_con_su_meta_no_solo_con_su_nombre(self):
        from tests.unit.test_compliance import full_session

        s = full_session()
        assert s.kpis and s.kpis[0].target, "el escenario no tiene meta que comprobar"
        texto = self._auditoria(s)
        assert s.kpis[0].name in texto
        assert s.kpis[0].target in texto, "el KPI llegó sin su meta"

    def test_la_nomenclatura_define_las_siglas_del_ffi(self):
        from rcm_runbook.models.session import FFIRegistro
        from tests.unit.test_compliance import full_session

        s = full_session()
        s.ffi_por_modo[next(iter(s.failure_modes))] = FFIRegistro(
            horas=1752.0, metodo="single_single", formula="FFI = 2 * Mtive * Mted / Mmf"
        )
        texto = self._auditoria(s)
        # TPEF aparecía tres veces en la propia fila que existe para glosar
        # siglas, sin glosarse.
        assert "TPEF = Tiempo Promedio Entre Fallas" in texto
        for sigla in ("Mtive =", "Mted =", "Mmf ="):
            assert sigla in texto, f"{sigla} no está definido en el entregable"

    def test_la_glosa_del_digest_corresponde_a_su_politica(self):
        # `any(n in texto for n in POLICY_LABELS_ES.values())` aceptaba cualquier
        # etiqueta: poner siempre la de MBC pasaba.
        from rcm_runbook.models.catalogs import POLICY_LABELS_ES
        from tests.unit.test_compliance import full_session

        s = full_session()
        texto = s.digest_es()
        for fmid, decision in s.decisions.items():
            nombre = POLICY_LABELS_ES[decision.policy]
            assert f"{decision.policy.value} ({nombre})" in texto, (
                f"{fmid} lleva la glosa de otra política"
            )

    def test_los_dos_avisos_del_motor_se_traducen(self):
        # El motor emite exactamente dos. La primera versión traducía uno y
        # tenía una tercera rama que no corresponde a ningún texto del motor:
        # código muerto que disfrazaba el hueco.
        import inspect

        from rcm_runbook.agent.tools import _aviso_en_espanol
        from rcm_runbook.engine import ffi as motor

        fuente = inspect.getsource(motor._build_result)
        emitidos = [linea for linea in fuente.splitlines() if "Computed FFI" in linea]
        assert len(emitidos) == 2, f"el motor cambió de avisos: {emitidos}"

        for crudo in (
            "Computed FFI (200000.0 h) exceeds the protective device MTBF "
            "(Mtive=1.0 h); the interval is suspect - review inputs.",
            "Computed FFI (0.0 h) is shorter than 24 h; failure finding this "
            "frequent is usually impractical - consider redesign.",
        ):
            traducido = _aviso_en_espanol(crudo)
            assert "Computed" not in traducido, f"sin traducir: {traducido[:60]}"
            assert "FFI calculado" in traducido

    def test_el_aviso_traducido_conserva_el_numero(self):
        from rcm_runbook.agent.tools import _aviso_en_espanol

        traducido = _aviso_en_espanol(
            "Computed FFI (0.0 h) is shorter than 24 h; failure finding this "
            "frequent is usually impractical - consider redesign."
        )
        assert "0.0" in traducido, "perdió el valor calculado, que es el dato útil"

    def test_al_completar_no_se_anuncia_una_septima_fase(self):
        from rcm_runbook.agent import tools as tools_mod
        from rcm_runbook.models.session import Phase, RCMSession

        sesion = RCMSession()
        sesion.phase = Phase.P6_PLAN
        avanzadas = tools_mod._avanzar_si_la_compuerta_esta_verde(sesion)
        if not avanzadas:  # la compuerta P6 no está verde en una sesión vacía
            avanzadas = ["7 (Análisis completado)"]
        tools_mod._ULTIMO_AVANCE["s-fin"] = avanzadas

        class Ctx:
            session_id = "s-fin"

        nota = tools_mod._nota_de_avance(Ctx())
        assert "fase 7" not in nota, "anuncia una séptima fase de seis"
        assert "COMPLETADO" in nota

    def test_advance_phase_tampoco_anuncia_la_fase_7(self):
        from rcm_runbook.agent import tools as tools_mod
        from tests.unit.test_compliance import full_session

        sesion = full_session()
        from rcm_runbook.models.session import Phase

        sesion.phase = Phase.P6_PLAN

        class Ctx:
            session_id = "s-adv"
            session_state = {"rcm": sesion.model_dump(mode="json")}

        salida = tools_mod.advance_phase.entrypoint(Ctx())
        assert "fase 7" not in salida, salida


class TestNingunaSiglaSeQuedaSinNombre:
    """Van trece en esta familia. Los diccionarios de etiquetas usan `.get(x,
    crudo)`: si mañana el motor añade un método de FFI o el dominio un tipo de
    función, vuelve el identificador inglés al entregable y nadie se entera. Es
    el mismo mecanismo que produjo «(proteccion)» sin tilde."""

    def test_toda_politica_tiene_nombre(self):
        from rcm_runbook.models.catalogs import POLICY_LABELS_ES, MaintenancePolicy

        faltan = [p.value for p in MaintenancePolicy if p not in POLICY_LABELS_ES]
        assert not faltan, f"políticas sin nombre en español: {faltan}"

    def test_todo_tipo_de_funcion_tiene_nombre(self):
        from rcm_runbook.models.catalogs import KIND_LABELS_ES
        from rcm_runbook.models.domain import FunctionKind

        faltan = [k.value for k in FunctionKind if k.value not in KIND_LABELS_ES]
        assert not faltan, f"tipos de función sin nombre: {faltan}"

    def test_todo_metodo_de_ffi_del_motor_tiene_nombre(self):
        from rcm_runbook.engine.ffi import _METHODS
        from rcm_runbook.models.catalogs import METODOS_FFI_ES

        faltan = [m for m in _METHODS if m not in METODOS_FFI_ES]
        assert not faltan, f"métodos del motor sin traducir: {faltan}"

    def test_toda_ruta_del_diagrama_tiene_significado(self):
        from rcm_runbook.models.catalogs import (
            ROUTE_LABELS_ES,
            EvidentRoute,
            HiddenRoute,
        )

        faltan = [
            r.value for r in list(EvidentRoute) + list(HiddenRoute)
            if r.value not in ROUTE_LABELS_ES
        ]
        assert not faltan, f"rutas sin significado: {faltan}"

    def test_la_politica_del_expediente_lleva_su_nombre(self):
        import tempfile

        from openpyxl import load_workbook

        from rcm_runbook.export.excel import export_xlsx
        from rcm_runbook.models.catalogs import POLICY_LABELS_ES
        from tests.unit.test_compliance import full_session

        s = full_session()
        with tempfile.TemporaryDirectory() as tmp:
            wb = load_workbook(export_xlsx(s, tmp))
        aud = " | ".join(
            str(c.value) for f in wb["AUDITORIA RCM"].iter_rows()
            for c in f if c.value is not None
        )
        for decision in s.decisions.values():
            nombre = POLICY_LABELS_ES[decision.policy]
            assert f"{decision.policy.value} ({nombre})" in aud, (
                f"{decision.policy.value} va pelado en el expediente"
            )

    def test_lookups_lleva_la_leyenda_de_politicas(self):
        # La etiqueta venía en el fixture del cliente y se estaba tirando.
        import tempfile

        from openpyxl import load_workbook

        from rcm_runbook.export.excel import export_xlsx
        from tests.unit.test_compliance import full_session

        with tempfile.TemporaryDirectory() as tmp:
            wb = load_workbook(export_xlsx(full_session(), tmp))
        texto = " | ".join(
            str(c.value) for f in wb["LOOKUPS"].iter_rows() for c in f if c.value
        )
        assert "Rd — " in texto and "ExEd — " in texto

    def test_los_codigos_iso_de_lookups_no_se_movieron_de_columna(self):
        # Las validaciones del libro apuntan a columnas concretas: añadir una
        # clave en medio del diccionario desplazó los códigos ISO y habría roto
        # los desplegables del cliente.
        import tempfile

        from openpyxl import load_workbook

        from rcm_runbook.export.excel import export_xlsx
        from rcm_runbook.models.catalogs import fixture
        from tests.unit.test_compliance import full_session

        with tempfile.TemporaryDirectory() as tmp:
            hoja = load_workbook(export_xlsx(full_session(), tmp))["LOOKUPS"]
        columna = [hoja.cell(row=r, column=5).value for r in range(2, 22)]
        esperado = [c.code for c in fixture().menu.iso14224_failure_mode_codes]
        assert columna == esperado


class TestElAvisoConservaLosDosNumeros:
    """`re.findall(...)[0]` se quedaba con la primera cifra de dos. En el chat el
    modelo reconstruía el Mtive del contexto; en la celda F de AUDITORIA no hay
    contexto que reconstruir. Y el test solo miraba uno de los dos avisos."""

    def test_el_aviso_del_mtbf_conserva_los_dos(self):
        from rcm_runbook.agent.tools import _aviso_en_espanol

        traducido = _aviso_en_espanol(
            "Computed FFI (306950.4 h) exceeds the protective device MTBF "
            "(Mtive=8760.0 h); the interval is suspect - review inputs."
        )
        assert "306950.4" in traducido, "perdió el FFI calculado"
        assert "8760.0" in traducido, "perdió el Mtive, que es con lo que se compara"

    def test_el_aviso_de_las_24_h_conserva_el_suyo(self):
        from rcm_runbook.agent.tools import _aviso_en_espanol

        traducido = _aviso_en_espanol(
            "Computed FFI (0.0 h) is shorter than 24 h; failure finding this "
            "frequent is usually impractical - consider redesign."
        )
        assert "0.0" in traducido

    def test_la_fila_de_nomenclatura_lleva_su_etiqueta(self):
        import tempfile

        from openpyxl import load_workbook

        from rcm_runbook.export.excel import export_xlsx
        from rcm_runbook.models.session import FFIRegistro
        from tests.unit.test_compliance import full_session

        s = full_session()
        s.ffi_por_modo[next(iter(s.failure_modes))] = FFIRegistro(
            horas=1752.0, metodo="single_single", formula="FFI = 2 * Mtive * Mted / Mmf"
        )
        with tempfile.TemporaryDirectory() as tmp:
            hoja = load_workbook(export_xlsx(s, tmp))["AUDITORIA RCM"]
        texto = " | ".join(
            str(c.value) for f in hoja.iter_rows() for c in f if c.value is not None
        )
        assert "Nomenclatura:" in texto, "la glosa quedó sin etiqueta que la anuncie"


class TestMarcarLaTareaDeVerdadCambiaElEstado:
    """El bloqueo mandaba llamar a `record_task(es_busqueda_de_fallas=True)`, la
    herramienta contestaba ✔, el estado no cambiaba y el entregable quedaba
    bloqueado para siempre: `campos` no incluía la marca, así que la llamada
    caía en la rama idempotente y la tiraba. Un rechazo sin salida es peor que
    no tener compuerta."""

    def _con_tarea(self, **extra):
        from rcm_runbook.models.domain import MaintenanceTask
        from tests.unit.test_compliance import full_session

        s = full_session()
        fm = next(f for f in s.failure_modes if s.effects[f].is_hidden)
        base = {
            "failure_mode_id": fm, "description": "Prueba funcional del disparo",
            "frequency": "Bimestral", "duration_hours": 4.0,
            "discipline": "Instrumentista",
        }
        s.tasks[fm] = [MaintenanceTask(**base)]
        return s, fm, MaintenanceTask(**base, **extra)

    def test_marcar_con_reemplazar_cambia_el_estado(self):
        s, fm, marcada = self._con_tarea(es_busqueda_de_fallas=True)
        s.add_task(marcada, reemplazar=True)
        assert s.tasks[fm][0].es_busqueda_de_fallas is True, "la marca se tiró"
        assert len(s.tasks[fm]) == 1

    def test_marcar_sin_reemplazar_avisa_en_vez_de_mentir(self):
        import pytest

        from rcm_runbook.errors import ReglaDeNegocio

        s, _, marcada = self._con_tarea(es_busqueda_de_fallas=True)
        with pytest.raises(ReglaDeNegocio, match="reemplazar=True"):
            s.add_task(marcada)

    def test_la_llamada_identica_sigue_siendo_idempotente(self):
        s, fm, igual = self._con_tarea()
        s.add_task(igual)
        assert len(s.tasks[fm]) == 1

    def test_el_bloqueo_se_puede_resolver(self):
        # El lazo cerrado completo: bloquea, se hace lo que dice el mensaje, y
        # deja de bloquear.
        from rcm_runbook.engine import compliance
        from rcm_runbook.models.session import FFIRegistro

        s, fm, marcada = self._con_tarea(es_busqueda_de_fallas=True)
        s.ffi_por_modo[fm] = FFIRegistro(
            horas=1752.0, metodo="single_single", formula="x"
        )
        assert [b for b in compliance.export_blockers(s) if "búsqueda de fallas" in b]
        s.add_task(marcada, reemplazar=True)
        assert not [
            b for b in compliance.export_blockers(s) if "búsqueda de fallas" in b
        ], "hizo lo que el mensaje pedía y el bloqueo siguió"


class TestUnFfiQueYaNoAplicaNoBloqueaParaSiempre:
    """`ffi_por_modo` es solo escritura: no hay herramienta que lo quite. Si el
    modo deja de ser oculto, pasa a operar-hasta-la-falla o se descarta por no
    creíble, el intervalo viejo bloqueaba el entregable sin salida — y
    `_gate_p6` salta esos mismos modos a propósito, así que `export_blockers` se
    contradecía consigo mismo."""

    def _sesion(self):
        from rcm_runbook.models.session import FFIRegistro
        from tests.unit.test_compliance import full_session

        s = full_session()
        fm = next(f for f in s.failure_modes if s.effects[f].is_hidden)
        s.ffi_por_modo[fm] = FFIRegistro(
            horas=1752.0, metodo="single_single", formula="x"
        )
        s.tasks[fm] = []
        return s, fm

    def _bloquea(self, s) -> bool:
        from rcm_runbook.engine import compliance

        return bool([b for b in compliance.export_blockers(s) if "búsqueda de fallas" in b])

    def test_un_modo_oculto_vigente_sin_tarea_sigue_bloqueando(self):
        s, _ = self._sesion()
        assert self._bloquea(s), "el contrapeso: aquí SÍ hace falta la tarea"

    def test_si_deja_de_ser_oculto_y_se_reevalua_ya_no_bloquea(self):
        # Corregir el efecto deja obsoleta la decisión; el escenario honesto es
        # corregir las dos cosas.
        s, fm = self._sesion()
        s.effects[fm] = s.effects[fm].model_copy(
            update={"is_hidden": False, "hidden_route": None}
        )
        s.decisions.pop(fm)
        assert not self._bloquea(s)

    def test_una_politica_de_busqueda_de_fallas_con_efecto_no_oculto_bloquea(self):
        # Contradicción del análisis: la política ES búsqueda de fallas. No es
        # razón para saltarse una compuerta de seguridad — las excepciones se
        # añadieron para no rechazar análisis correctos, y este no lo es.
        s, fm = self._sesion()
        s.effects[fm] = s.effects[fm].model_copy(
            update={"is_hidden": False, "hidden_route": None}
        )
        assert self._bloquea(s)

    def test_un_modo_oculto_sin_decision_todavia_bloquea(self):
        s, fm = self._sesion()
        s.decisions.pop(fm)
        assert self._bloquea(s)

    def test_un_modo_sin_efecto_registrado_bloquea(self):
        # Falta información, que no es lo mismo que estar bien.
        s, fm = self._sesion()
        s.effects.pop(fm)
        s.decisions.pop(fm)
        assert self._bloquea(s)

    def test_si_la_politica_pasa_a_operar_hasta_la_falla_ya_no_bloquea(self):
        from rcm_runbook.models.catalogs import MaintenancePolicy

        s, fm = self._sesion()
        s.decisions[fm] = s.decisions[fm].model_copy(
            update={"policy": MaintenancePolicy.OHF}
        )
        assert not self._bloquea(s)

    def test_si_el_modo_se_descarta_por_no_creible_ya_no_bloquea(self):
        s, fm = self._sesion()
        s.failure_modes[fm] = s.failure_modes[fm].model_copy(
            update={"credible": False,
                    "non_credible_discard": "No aplica en este contexto operacional"}
        )
        assert not self._bloquea(s)


class TestElEntregableEnsenaQuienEjecutaElFfi:
    """Quién marca la fila es el modelo, y ninguna comprobación determinista
    puede saber si marcó la correcta: se vio marcar una limpieza diaria dejando
    la prueba funcional real a «Parada de Planta». Lo que sí se puede es ponerlo
    delante de quien revisa."""

    def _auditoria(self, s) -> str:
        import tempfile

        from openpyxl import load_workbook

        from rcm_runbook.export.excel import export_xlsx

        with tempfile.TemporaryDirectory() as tmp:
            hoja = load_workbook(export_xlsx(s, tmp))["AUDITORIA RCM"]
        return " | ".join(
            str(c.value) for f in hoja.iter_rows() for c in f if c.value is not None
        )

    def test_dice_que_tarea_ejecuta_el_intervalo_y_cada_cuanto(self):
        from rcm_runbook.models.domain import MaintenanceTask
        from rcm_runbook.models.session import FFIRegistro
        from tests.unit.test_compliance import full_session

        s = full_session()
        fm = next(f for f in s.failure_modes if s.effects[f].is_hidden)
        s.ffi_por_modo[fm] = FFIRegistro(horas=1752.0, metodo="single_single", formula="x")
        s.tasks[fm] = [MaintenanceTask(
            failure_mode_id=fm, description="Prueba funcional del disparo",
            frequency="Bimestral", duration_hours=4.0, discipline="Instrumentista",
            es_busqueda_de_fallas=True,
        )]
        texto = self._auditoria(s)
        assert "Ejecutado por: Prueba funcional del disparo (Bimestral)" in texto

    def test_lo_dice_tambien_cuando_no_hay_ninguna_marcada(self):
        from rcm_runbook.models.session import FFIRegistro
        from tests.unit.test_compliance import full_session

        s = full_session()
        fm = next(f for f in s.failure_modes if s.effects[f].is_hidden)
        s.ffi_por_modo[fm] = FFIRegistro(horas=1752.0, metodo="single_single", formula="x")
        s.tasks[fm] = []
        assert "NINGUNA TAREA MARCADA" in self._auditoria(s)


class TestElMensajeDeBloqueoDiceQueArreglarYDonde:
    """Tres mutantes sobrevivían: quitar el nombre del modo, quitar la lista de
    tareas registradas, y forzar la marca a False al reemplazar. Un bloqueo que
    no dice qué modo ni qué hay registrado deja al operador sin por dónde
    empezar."""

    def _bloqueo(self) -> str:
        from rcm_runbook.engine import compliance
        from rcm_runbook.models.domain import MaintenanceTask
        from rcm_runbook.models.session import FFIRegistro
        from tests.unit.test_compliance import full_session

        s = full_session()
        fm = next(f for f in s.failure_modes if s.effects[f].is_hidden)
        s.ffi_por_modo[fm] = FFIRegistro(horas=1752.0, metodo="single_single", formula="x")
        s.tasks[fm] = [MaintenanceTask(
            failure_mode_id=fm, description="Limpieza del entorno", frequency="Mensual",
            duration_hours=1.0, discipline="Mecánico",
        )]
        self.fm = fm
        return next(b for b in compliance.export_blockers(s) if "búsqueda de fallas" in b)

    def test_nombra_el_modo(self):
        bloqueo = self._bloqueo()
        assert self.fm in bloqueo, "el operador no sabe qué modo arreglar"

    def test_lista_las_tareas_ya_registradas(self):
        assert "Limpieza del entorno" in self._bloqueo()

    def test_dice_como_arreglarlo(self):
        assert "es_busqueda_de_fallas=True" in self._bloqueo()


class TestLaLeyendaEstaDondeSeVe:
    """La leyenda de políticas se puso en LOOKUPS, que es una hoja oculta: el
    comentario decía «quien abre el libro veía Rd y ExEd sin nada que los
    explicara en ninguna hoja» y seguía sin verlo. Un arreglo invisible no es un
    arreglo, y su test pasaba igual."""

    def _libro(self):
        import tempfile

        from openpyxl import load_workbook

        from rcm_runbook.export.excel import export_xlsx
        from tests.unit.test_compliance import full_session

        with tempfile.TemporaryDirectory() as tmp:
            return load_workbook(export_xlsx(full_session(), tmp))

    def test_lookups_sigue_oculta(self):
        assert self._libro()["LOOKUPS"].sheet_state == "hidden"

    def test_las_politicas_se_explican_en_una_hoja_visible(self):
        from rcm_runbook.models.catalogs import POLICY_LABELS_ES

        wb = self._libro()
        visibles = [n for n in wb.sheetnames if wb[n].sheet_state == "visible"]
        texto = " | ".join(
            str(c.value) for n in visibles for f in wb[n].iter_rows()
            for c in f if c.value is not None
        )
        for politica, nombre in POLICY_LABELS_ES.items():
            assert nombre in texto, f"{politica.value} sin explicar en hoja visible"


class TestLaDisposicionDeLaHojaDeAuditoria:
    """Todos los tests de esta hoja concatenaban las filas en un string y
    buscaban subcadenas, así que fila y columna eran libres: la leyenda de
    políticas aterrizó ENTRE la cabecera del FFI y su tabla —doce filas de por
    medio, leyéndose como si MBC y MBT fueran modos de falla— con su test en
    verde. Un arreglo colocado donde no toca sigue sin arreglar nada."""

    def _hoja(self):
        import tempfile

        from openpyxl import load_workbook

        from rcm_runbook.export.excel import export_xlsx
        from rcm_runbook.models.session import FFIRegistro
        from tests.unit.test_compliance import full_session

        s = full_session()
        fm = next(f for f in s.failure_modes if s.effects[f].is_hidden)
        s.ffi_por_modo[fm] = FFIRegistro(
            horas=1752.0, metodo="single_single", formula="FFI = 2 * Mtive * Mted / Mmf"
        )
        with tempfile.TemporaryDirectory() as tmp:
            return load_workbook(export_xlsx(s, tmp))["AUDITORIA RCM"], fm

    def _fila_de(self, hoja, texto: str, desde: int = 1) -> int:
        for fila in hoja.iter_rows(min_row=desde):
            for celda in fila:
                if celda.value and str(celda.value).startswith(texto):
                    return celda.row
        raise AssertionError(f"no encontré «{texto}» desde la fila {desde}")

    def _fila_del_ffi(self, hoja, fm: str) -> int:
        # Desde la cabecera: el id del modo aparece antes en HITL y en las
        # justificaciones, y buscarlo desde arriba daba la fila equivocada.
        return self._fila_de(hoja, fm, desde=self._fila_de(hoja, "Intervalo de búsqueda"))

    def test_la_tabla_del_ffi_va_pegada_a_su_cabecera(self):
        hoja, fm = self._hoja()
        cabecera = self._fila_de(hoja, "Intervalo de búsqueda de fallas")
        datos = self._fila_del_ffi(hoja, fm)
        assert 0 < datos - cabecera <= 3, (
            f"hay {datos - cabecera - 1} filas entre la cabecera del FFI y su tabla"
        )

    def test_la_leyenda_de_politicas_no_esta_dentro_de_la_seccion_del_ffi(self):
        hoja, fm = self._hoja()
        leyenda = self._fila_de(hoja, "Políticas de mantenimiento")
        datos_ffi = self._fila_del_ffi(hoja, fm)
        assert leyenda > datos_ffi, (
            "la leyenda quedó entre la cabecera del FFI y sus datos: bajo esa "
            "cabecera, MBC y MBT se leen como modos de falla"
        )

    def test_ejecutado_por_va_en_la_columna_del_ffi(self):
        # Moverla de columna no rompía ningún test, y en una hoja ancha eso la
        # deja lejos del intervalo al que se refiere.
        hoja, fm = self._hoja()
        fila = self._fila_del_ffi(hoja, fm)
        valores = {c.column_letter: str(c.value) for c in hoja[fila] if c.value}
        assert "F" in valores and valores["F"].startswith("Ejecutado por:"), (
            f"«Ejecutado por» no está en la columna F: {valores}"
        )

    def test_cada_seccion_aparece_una_sola_vez(self):
        hoja, _ = self._hoja()
        titulos = [
            str(c.value) for fila in hoja.iter_rows() for c in fila
            if c.value and c.font and c.font.bold
        ]
        assert len(titulos) == len(set(titulos)), f"secciones duplicadas: {titulos}"


class TestElRechazoDiceQueCampoCambia:
    """El mensaje decía «otros datos» y a continuación imprimía datos idénticos,
    porque el único campo que cambiaba —la marca de búsqueda de fallas— no salía
    en él. Y es justo la llamada que el bloqueo del FFI ordena hacer: el camino
    documentado como correcto terminaba en un rechazo visiblemente falso."""

    def _chocar(self, **extra):
        from rcm_runbook.errors import ReglaDeNegocio
        from rcm_runbook.models.domain import MaintenanceTask
        from tests.unit.test_compliance import full_session

        s = full_session()
        fm = next(f for f in s.failure_modes if s.effects[f].is_hidden)
        base = {
            "failure_mode_id": fm, "description": "Prueba funcional del disparo",
            "frequency": "Bimestral", "duration_hours": 4.0,
            "discipline": "Instrumentista",
        }
        s.tasks[fm] = [MaintenanceTask(**base)]
        try:
            s.add_task(MaintenanceTask(**{**base, **extra}))
        except ReglaDeNegocio as exc:
            return str(exc)
        raise AssertionError("no rechazó")

    def test_al_marcar_el_mensaje_dice_que_la_marca_es_lo_que_cambia(self):
        mensaje = self._chocar(es_busqueda_de_fallas=True)
        assert "sin marcar como búsqueda de fallas" in mensaje
        assert "marcada como búsqueda de fallas" in mensaje

    def test_las_dos_fichas_no_son_identicas(self):
        # El síntoma exacto: «otros datos» seguido de dos líneas iguales.
        mensaje = self._chocar(es_busqueda_de_fallas=True)
        fichas = [
            ln for ln in mensaje.splitlines()
            if ln.strip().startswith(("Registrada:", "La nueva:"))
        ]
        assert len(fichas) == 2
        assert fichas[0].split("(", 1)[1] != fichas[1].split("(", 1)[1], (
            f"el rechazo imprime los mismos datos dos veces: {fichas}"
        )

    def test_el_casi_duplicado_tambien_lo_dice(self):
        mensaje = self._chocar(
            description="Prueba funcional de disparo", es_busqueda_de_fallas=True
        )
        assert "casi idéntica" in mensaje
        assert "marcada como búsqueda de fallas" in mensaje


class TestElFfiSoloAplicaAFallasOcultas:
    """La herramienta contestaba «queda registrado y sale en el entregable» para
    un modo evidente, la compuerta saltaba ese modo por no ser oculto, y el Excel
    salía con un intervalo que no gobierna nada. La promesa de la herramienta,
    desmentida en silencio por la compuerta."""

    def _ctx(self):
        from tests.unit.test_compliance import full_session

        sesion = full_session()

        class Ctx:
            session_id = "s-ffi-oculto"
            session_state = {"rcm": sesion.model_dump(mode="json")}

        return Ctx(), sesion

    def test_un_modo_evidente_se_rechaza(self):
        from rcm_runbook.agent import tools as tools_mod

        ctx, sesion = self._ctx()
        evidente = next(f for f in sesion.failure_modes if not sesion.effects[f].is_hidden)
        salida = tools_mod.calculate_ffi.entrypoint(
            ctx, method="single_single", mtive_hours=43800, mted_hours=17520,
            mmf_hours=876000, failure_mode_id=evidente,
        )
        assert "no es una falla oculta" in salida
        assert "sale en el entregable" not in salida, "prometió lo que no cumple"

    def test_un_modo_oculto_se_registra(self):
        from rcm_runbook.agent import tools as tools_mod
        from rcm_runbook.models.session import RCMSession

        ctx, sesion = self._ctx()
        oculto = next(f for f in sesion.failure_modes if sesion.effects[f].is_hidden)
        salida = tools_mod.calculate_ffi.entrypoint(
            ctx, method="single_single", mtive_hours=43800, mted_hours=17520,
            mmf_hours=876000, failure_mode_id=oculto,
        )
        assert "sale en el entregable" in salida
        guardada = RCMSession.model_validate(ctx.session_state["rcm"])
        assert oculto in guardada.ffi_por_modo


class TestElFfiNoSobreviveASuModo:
    """El intervalo se calculaba antes de registrar el efecto, el efecto
    resultaba evidente, y nadie retiraba el FFI: viajaba al Excel con «NINGUNA
    TAREA MARCADA» mientras la compuerta saltaba ese modo por no ser oculto.
    Descarte silencioso número nueve, y el mismo que el commit anterior decía
    haber cerrado — reabierto por el orden inverso de llamadas."""

    def _ctx(self, quitar_efecto: str = ""):
        from tests.unit.test_compliance import full_session

        sesion = full_session()
        fm = next(f for f in sesion.failure_modes if sesion.effects[f].is_hidden)
        if quitar_efecto:
            sesion.effects.pop(fm)

        class Ctx:
            session_id = "s-ffi-orden"
            session_state = {"rcm": sesion.model_dump(mode="json")}

        return Ctx(), fm

    def _calcular(self, ctx, fm):
        from rcm_runbook.agent import tools as tools_mod

        return tools_mod.calculate_ffi.entrypoint(
            ctx, method="single_single", mtive_hours=43800, mted_hours=17520,
            mmf_hours=876000, failure_mode_id=fm,
        )

    def test_sin_efecto_registrado_no_se_calcula(self):
        ctx, fm = self._ctx(quitar_efecto="sí")
        salida = self._calcular(ctx, fm)
        assert "no tiene efectos registrados" in salida
        assert "sale en el entregable" not in salida

    def test_corregir_el_efecto_a_evidente_saca_el_ffi_del_entregable(self):
        # NO se borra el dato: borrarlo dejaba a la compuerta sin nada que
        # visitar y desactivaba la excepción de política BF, así que un modo de
        # seguridad sin tarea que lo ejecute salía limpio. Se filtra al exportar,
        # con el mismo criterio que usa la compuerta.
        import tempfile

        from openpyxl import load_workbook

        from rcm_runbook.export.excel import export_xlsx
        from rcm_runbook.models.session import RCMSession

        ctx, fm = self._ctx()
        self._calcular(ctx, fm)
        sesion = RCMSession.model_validate(ctx.session_state["rcm"])
        sesion.decisions.pop(fm, None)  # reevaluada tras corregir el efecto
        sesion.set_effect(
            fm, local="Se detiene la bomba sin aviso previo",
            system="Pérdida de caudal en la línea", plant="Parada de la unidad",
            is_hidden=False, operational=True,
        )
        assert fm in sesion.ffi_por_modo, "el dato no se tira, solo deja de aplicar"
        with tempfile.TemporaryDirectory() as tmp:
            hoja = load_workbook(export_xlsx(sesion, tmp))["AUDITORIA RCM"]
        texto = " | ".join(
            str(c.value) for f in hoja.iter_rows() for c in f if c.value is not None
        )
        assert "Ejecutado por" not in texto, "el FFI que ya no aplica salió al Excel"

    def test_si_la_politica_sigue_siendo_busqueda_de_fallas_la_compuerta_no_se_apaga(self):
        # El caso que rompía el borrado: efecto corregido a evidente pero la
        # decisión sigue siendo BF. Es una contradicción del análisis y no puede
        # saltarse una compuerta de seguridad.
        from rcm_runbook.engine import compliance
        from rcm_runbook.models.session import RCMSession

        ctx, fm = self._ctx()
        self._calcular(ctx, fm)
        sesion = RCMSession.model_validate(ctx.session_state["rcm"])
        sesion.set_effect(
            fm, local="Se detiene la bomba sin aviso previo",
            system="Pérdida de caudal en la línea", plant="Parada de la unidad",
            is_hidden=False, operational=True,
        )
        bloqueos = [
            b for b in compliance.export_blockers(sesion) if "búsqueda de fallas" in b
        ]
        assert bloqueos, "un modo con política BF y sin tarea salió sin bloquear"

    def test_corregir_el_efecto_manteniendolo_oculto_lo_conserva(self):
        # El contrapeso: reformular el efecto sin cambiar su naturaleza no puede
        # tirar un cálculo válido.
        from rcm_runbook.models.session import RCMSession

        ctx, fm = self._ctx()
        self._calcular(ctx, fm)
        sesion = RCMSession.model_validate(ctx.session_state["rcm"])
        previo = sesion.effects[fm]
        sesion.set_effect(
            fm, local=previo.local, system=previo.system,
            plant="Redacción corregida del efecto en planta",
            is_hidden=True, safety=previo.safety, environment=previo.environment,
            operational=previo.operational, non_operational=previo.non_operational,
        )
        assert fm in sesion.ffi_por_modo


class TestElFfiCaducadoSeAnotaNoSeEsconde:
    """Filtrar del entregable los FFI que ya no aplican dejaba la sección con su
    cabecera, su nomenclatura y cero filas —definiciones sin término— y hacía
    desaparecer un intervalo que sí se calculó, sin decirlo. Es el mismo
    descarte silencioso un piso más arriba."""

    def _hoja(self, vigente: bool):
        import tempfile

        from openpyxl import load_workbook

        from rcm_runbook.export.excel import export_xlsx
        from rcm_runbook.models.session import FFIRegistro
        from tests.unit.test_compliance import full_session

        s = full_session()
        fm = next(f for f in s.failure_modes if s.effects[f].is_hidden)
        s.ffi_por_modo[fm] = FFIRegistro(
            horas=1752.0, metodo="single_single", formula="FFI = 2 * Mtive * Mted / Mmf"
        )
        if not vigente:
            s.decisions.pop(fm)
            s.set_effect(
                fm, local="Se detiene la bomba sin aviso previo",
                system="Pérdida de caudal en la línea", plant="Parada de la unidad",
                is_hidden=False, operational=True,
            )
        with tempfile.TemporaryDirectory() as tmp:
            hoja = load_workbook(export_xlsx(s, tmp))["AUDITORIA RCM"]
        return " | ".join(
            str(c.value) for f in hoja.iter_rows() for c in f if c.value is not None
        ), fm

    def test_el_intervalo_caducado_sigue_en_el_entregable(self):
        texto, fm = self._hoja(vigente=False)
        assert "1752" in texto, "el intervalo calculado desapareció sin decirlo"
        assert fm in texto

    def test_y_dice_que_ya_no_aplica(self):
        texto, _ = self._hoja(vigente=False)
        assert "YA NO APLICA" in texto
        assert "Ejecutado por" not in texto

    def test_el_vigente_dice_quien_lo_ejecuta(self):
        texto, _ = self._hoja(vigente=True)
        assert "Ejecutado por" in texto
        assert "YA NO APLICA" not in texto

    def test_la_nomenclatura_nunca_queda_sin_filas(self):
        # Definiciones sin término: la cabecera y la glosa escritas, y ninguna
        # fila debajo.
        texto, fm = self._hoja(vigente=False)
        assert "Nomenclatura:" in texto and fm in texto


class TestElFalloDelProveedorNoSaleEnInglesNiConLaFacturacion:
    """Criterio 44 del UAT, y 49. El fallo del proveedor NO llega como error
    HTTP: llega con 200 y el inglés dentro del `content` del turno, con el
    estado de facturación del operador. Estaba traducido solo en `demo.html`,
    o sea en el navegador: la UI de os.agno.com y cualquier cliente de la API
    seguían viendo «Your credit balance is too low… Plans & Billing».
    """

    CRUDO = (
        "Error code: 400 - {'type': 'error', 'error': {'type': "
        "'invalid_request_error', 'message': 'Your credit balance is too low to "
        "access the Anthropic API. Please go to Plans & Billing to upgrade or "
        "purchase credits.'}, 'request_id': 'req_011CdcXjgktwbRwnTeM3XbG5'}"
    )

    def test_traduce_y_no_deja_ni_una_palabra_de_la_facturacion(self):
        from rcm_runbook.app import en_espanol_si_es_fallo_del_proveedor

        salida = en_espanol_si_es_fallo_del_proveedor(self.CRUDO)
        for prohibido in ("credit balance", "Plans & Billing", "upgrade",
                          "purchase credits", "Please", "Anthropic API",
                          "request_id", "invalid_request_error"):
            assert prohibido not in salida, f"se filtró «{prohibido}»"
        assert "no está disponible" in salida

    @pytest.mark.parametrize(
        "crudo",
        [
            "Error code: 429 - rate_limit_error: too many requests",
            "Error code: 529 - {'type': 'error', 'error': {'type': 'overloaded_error'}}",
            "Error code: 401 - authentication_error: invalid x-api-key",
        ],
    )
    def test_los_demas_fallos_del_proveedor_tambien(self, crudo):
        from rcm_runbook.app import en_espanol_si_es_fallo_del_proveedor

        salida = en_espanol_si_es_fallo_del_proveedor(crudo)
        assert salida != crudo
        assert not re.search(r"[a-z]+_error|Error code", salida)

    @pytest.mark.parametrize(
        "legitima",
        [
            "El modo FTS es el más frecuente en bombas centrífugas.",
            "Hubo un error al registrar la tarea: la frecuencia es obligatoria.",
            "Le devuelvo el borrador del Excel con lo registrado hasta ahora.",
        ],
    )
    def test_una_respuesta_legitima_no_se_convierte_en_un_aviso_de_averia(
        self, legitima
    ):
        # El riesgo simétrico: traducir de más convierte una respuesta buena del
        # agente en «el servicio no está disponible», que es una avería falsa.
        from rcm_runbook.app import en_espanol_si_es_fallo_del_proveedor

        assert en_espanol_si_es_fallo_del_proveedor(legitima) == legitima

    def test_el_turno_completo_sale_traducido_por_la_api_no_por_el_navegador(
        self, client
    ):
        """Extremo a extremo por el middleware: es lo que distingue este arreglo
        del que ya había. Si solo tradujera el navegador, esto seguiría en
        inglés."""
        from rcm_runbook import app as app_mod

        async def responder(scope, receive, send):
            cuerpo = json.dumps(
                {"content": self.CRUDO, "session_id": "s-fallo"}
            ).encode()
            await send({
                "type": "http.response.start", "status": 200,
                "headers": [(b"content-type", b"application/json"),
                            (b"content-length", str(len(cuerpo)).encode())],
            })
            await send({"type": "http.response.body", "body": cuerpo})

        mw = app_mod.TraducirFallosDelProveedor(responder)
        recibido: list[dict] = []

        async def recoger(mensaje):
            recibido.append(mensaje)

        async def vacio():
            return {"type": "http.request", "body": b"", "more_body": False}

        asyncio.run(mw({"type": "http", "path": "/agents/x/runs"}, vacio, recoger))
        cuerpo = json.loads(
            b"".join(m.get("body", b"") for m in recibido if m["type"].endswith("body"))
        )
        assert "credit balance" not in cuerpo["content"]
        assert "no está disponible" in cuerpo["content"]
        assert cuerpo["session_id"] == "s-fallo", "se comió el resto de la respuesta"
        largo = dict(
            (k.decode().lower(), v.decode())
            for m in recibido if m["type"].endswith("start") for k, v in m["headers"]
        )["content-length"]
        assert int(largo) == len(
            b"".join(m.get("body", b"") for m in recibido if m["type"].endswith("body"))
        ), "el content-length quedó mintiendo tras reescribir"

    def test_lo_que_no_es_json_pasa_intacto(self):
        """El streaming de os.agno.com va por SSE y NO se toca: bufferearlo
        rompería el backpressure que consume su UI. Queda descubierto y
        anotado, no resuelto: por ahí el inglés todavía sale."""
        from rcm_runbook import app as app_mod

        crudo = b"data: " + self.CRUDO.encode()

        async def responder(scope, receive, send):
            await send({
                "type": "http.response.start", "status": 200,
                "headers": [(b"content-type", b"text/event-stream")],
            })
            await send({"type": "http.response.body", "body": crudo})

        recibido: list[dict] = []

        async def recoger(mensaje):
            recibido.append(mensaje)

        async def vacio():
            return {"type": "http.request", "body": b"", "more_body": False}

        asyncio.run(
            app_mod.TraducirFallosDelProveedor(responder)(
                {"type": "http", "path": "/agents/x/runs"}, vacio, recoger
            )
        )
        assert any(m.get("body") == crudo for m in recibido)
        # Y la cabecera tiene que salir: sin ella la respuesta nunca arranca.
        arranques = [m for m in recibido if m["type"] == "http.response.start"]
        assert len(arranques) == 1, f"cabeceras enviadas: {len(arranques)}"
        assert dict(arranques[0]["headers"])[b"content-type"] == b"text/event-stream"
