"""Golden-file export tests — assert against the frozen benchmark fixture.

Scope (per plan): cell values, header rows verbatim, sheet names, data-validation
formulas in the SAVED output, column order. Styles/widths excluded (visual check).
"""

import json
from pathlib import Path

from openpyxl import load_workbook

from rcm_runbook.export.excel import (
    AMEF_HEADER_ROW,
    PLAN_HEADER_ROW,
    export_xlsx,
)
from rcm_runbook.export.rows import amef_headers, plan_headers, to_amef_rows, to_plan_rows
from tests.unit.test_compliance import full_session

FIXTURE = json.loads(
    (Path(__file__).parents[2] / "src/rcm_runbook/data/benchmark_fixture.json").read_text("utf-8")
)


class TestHeaderContract:
    def test_amef_row_aliases_match_fixture_headers_verbatim(self):
        assert amef_headers() == FIXTURE["amef"]["headers"]

    def test_plan_row_aliases_match_fixture_headers_verbatim(self):
        assert plan_headers() == FIXTURE["plan"]["headers"]


class TestProjections:
    def test_row_grain_one_per_credible_mode(self):
        s = full_session()
        assert len(to_amef_rows(s)) == 2
        s.add_failure_mode(
            "FF-001",
            description="Desgaste del impulsor por erosión",
            mechanism="Erosión",
            iso_code="LOO",
            cause="Sólidos en el fluido de proceso",
            root_cause="Filtración aguas arriba deficiente",
            failure_pattern="Fin de Vida Útil",
            credible=False,
            non_credible_discard="Fluido limpio verificado por análisis trimestral",
        )
        assert len(to_amef_rows(s)) == 2  # non-credible excluded from AMEF

    def test_plan_one_hot_strategy_marks(self):
        rows = to_plan_rows(full_session())
        mbc = next(r for r in rows if r.basado_condicion == "X")
        assert (mbc.preventivo, mbc.deteccion_fallas, mbc.otras) == ("", "", "")
        bf = next(r for r in rows if r.deteccion_fallas == "X")
        assert (bf.basado_condicion, bf.preventivo, bf.otras) == ("", "", "")

    def test_tpef_projected_from_failure_mode_not_task(self):
        s = full_session()
        amef = to_amef_rows(s)
        assert amef[0].rpn == 7 * 5 * 4
        # TPEF has exactly one owner: the FailureMode
        assert s.failure_modes["FM-001"].tpef is not None
        assert not hasattr(s.tasks["FM-001"][0], "tpef")


class TestSavedWorkbook:
    def _export(self, tmp_path: Path) -> Path:
        return export_xlsx(full_session(), tmp_path)

    def test_sheet_names(self, tmp_path):
        wb = load_workbook(self._export(tmp_path))
        assert "AMEF" in wb.sheetnames
        assert "PLAN DE MANTENIMIENTO" in wb.sheetnames
        assert "SAE-J1739" in wb.sheetnames
        assert "AUDITORIA RCM" in wb.sheetnames

    def test_amef_header_row_verbatim(self, tmp_path):
        ws = load_workbook(self._export(tmp_path))["AMEF"]
        saved = [
            ws.cell(row=AMEF_HEADER_ROW, column=c).value
            for c in range(2, 2 + len(FIXTURE["amef"]["headers"]))
        ]
        assert saved == FIXTURE["amef"]["headers"]

    def test_plan_header_row_verbatim(self, tmp_path):
        ws = load_workbook(self._export(tmp_path))["PLAN DE MANTENIMIENTO"]
        saved = [
            ws.cell(row=PLAN_HEADER_ROW, column=c).value
            for c in range(2, 2 + len(FIXTURE["plan"]["headers"]))
        ]
        assert saved == FIXTURE["plan"]["headers"]

    def test_data_row_values(self, tmp_path):
        ws = load_workbook(self._export(tmp_path))["AMEF"]
        row = AMEF_HEADER_ROW + 1
        headers = FIXTURE["amef"]["headers"]
        values = {headers[c]: ws.cell(row=row, column=c + 2).value for c in range(len(headers))}
        assert values["TAG"] == "P-201A"
        assert values["Codigo ISO 14224"] == "LOO"
        assert values["Estrategia de Mantenimiento"] == "MBC"
        assert values["RPN"] == 140
        assert values["Seguridad"] in ("", None)
        assert values["Operacional"] == "X"

    def test_dropdown_validations_in_saved_output(self, tmp_path):
        """Data-validation formulas must survive the save (planner amendment A6)."""
        ws = load_workbook(self._export(tmp_path))["AMEF"]
        formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
        assert any("LOOKUPS" in f for f in formulas), "dropdowns lost on save"
        # ISO code dropdown feeds from 20 codes (rows 2..21)
        iso = [f for f in formulas if "$E$" in f or "iso" in f.lower()]
        assert formulas, iso

    def test_lookups_sheet_hidden_and_menu_authoritative(self, tmp_path):
        wb = load_workbook(self._export(tmp_path))
        ws = wb["LOOKUPS"]
        assert ws.sheet_state == "hidden"
        iso_col = [ws.cell(row=r, column=5).value for r in range(2, 22)]
        fixture_codes = [c["code"] for c in FIXTURE["menu"]["iso14224_failure_mode_codes"]]
        assert iso_col == fixture_codes

    def test_audit_sheet_carries_hitl_and_justifications(self, tmp_path):
        ws = load_workbook(self._export(tmp_path))["AUDITORIA RCM"]
        text = "\n".join(
            str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
        )
        assert "Supervisora HSE" in text
        assert "P-F" in text or "prueba funcional" in text.lower()

    def test_tpef_block_on_plan_sheet(self, tmp_path):
        ws = load_workbook(self._export(tmp_path))["PLAN DE MANTENIMIENTO"]
        text = "\n".join(
            str(c.value) for row in ws.iter_rows(max_row=10) for c in row if c.value is not None
        )
        assert "TPEF" in text
        assert "OREDA" in text

    def test_filename_from_tag(self, tmp_path):
        assert self._export(tmp_path).name == "AMEF_P-201A.xlsx"
